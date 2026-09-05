# -*- coding: utf-8 -*-
"""
第四稿：让学生审辩AI的证明——依托"三个助手"、指向推理能力的"命题与证明"习题讲评课
落法：AI(备课助手)真实生成含漏洞证明 → 全班投屏实时审辩 → 三个助手采集每个学生的判断(过程性数据)
输出：/Users/chengqp/Desktop/tan/AI教证明_以教为学_论文.docx
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

BASE = "/Users/chengqp/Desktop/tan"
IMG = os.path.join(BASE, "图")
os.makedirs(IMG, exist_ok=True)
OUT = os.path.join(BASE, "AI教证明_以教为学_论文.docx")
SONG, HEI, KAI = "宋体", "黑体", "楷体"
FONT_PATH = "/System/Library/Fonts/Hiragino Sans GB.ttc"
def F(size): return ImageFont.truetype(FONT_PATH, size)

def set_cn_font(run, cn=SONG, size=10.5, bold=False, color=None):
    run.font.name = cn
    run._element.rPr.rFonts.set(qn('w:eastAsia'), cn)
    run.font.size = Pt(size); run.font.bold = bold
    if color: run.font.color.rgb = color

def add_para(doc, indent_chars=2, space_after=0, line=None, align=None):
    p = doc.add_paragraph(); pf = p.paragraph_format
    if indent_chars: pf.first_line_indent = Pt(10.5 * indent_chars)
    pf.space_after = Pt(space_after); pf.space_before = Pt(0)
    if line: pf.line_spacing = line; pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if align: p.alignment = align
    return p

def runs(p, segments, size=10.5):
    for text, style in segments:
        r = p.add_run(text); set_cn_font(r, SONG, size)
        if style == 'i': r.font.italic = True; r.font.name = 'Times New Roman'
        elif style == 'b': r.font.bold = True
        elif style == 'sup': r.font.superscript = True
    return p

def heading(doc, text, size=12, bold=True, before=8, after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text); set_cn_font(r, HEI, size, bold=bold)
    return p
print("样式就绪")


# ---------- 图1：审辩路径（全班共审·个体留痕） ----------
def draw_center_text(d, box, text, font, fill=(40,40,40)):
    x0,y0,x1,y1 = box; lines = text.split("\n")
    hs = [d.textbbox((0,0),ln,font=font)[3]-d.textbbox((0,0),ln,font=font)[1] for ln in lines]
    total = sum(hs)+(len(lines)-1)*6; cy=(y0+y1)/2-total/2
    for ln,h in zip(lines,hs):
        w=d.textbbox((0,0),ln,font=font)[2]-d.textbbox((0,0),ln,font=font)[0]
        d.text(((x0+x1)/2-w/2,cy),ln,font=font,fill=fill); cy+=h+6

def make_fig1():
    W,H=1580,580; img=Image.new("RGB",(W,H),"white"); d=ImageDraw.Draw(img)
    d.text((40,22),"“问‘凭什么’”审辩路径：全班实时共审一个会犯错的 AI · 每人的判断被平台留痕",font=F(27),fill=(15,15,15))
    boxes=[("① 立徒弟","约定当审稿人\n把 AI 证明每一步\n都追问到“有据”"),
           ("② 判真假","投屏 AI 判假命题\n每人互动作答\n举反例，平台统计"),
           ("③ 揪漏洞·补依据","全班逼问投屏 AI\n每人标出没依据的步\n平台统计漏洞分布"),
           ("④ 辨原逆","AI 以为逆命题必真\n学生举反例\n辨析原、逆命题"),
           ("⑤ 独立证·提交","独立证同类变式\n平台提交+同伴互评\n典型再交 AI 挑刺")]
    n=len(boxes); margin=40; gap=22; bw=(W-2*margin-(n-1)*gap)/n; top=140; bh=270
    colors=[(232,244,253),(255,243,224),(232,245,233),(243,229,245),(224,247,250)]
    line_c=[(21,101,192),(230,126,34),(56,142,60),(123,31,162),(0,131,143)]
    for i,(title,body) in enumerate(boxes):
        x0=margin+i*(bw+gap); x1=x0+bw; y0,y1=top,top+bh
        d.rounded_rectangle([x0,y0,x1,y1],radius=16,fill=colors[i],outline=line_c[i],width=3)
        d.rounded_rectangle([x0,y0,x1,y0+46],radius=16,fill=line_c[i])
        tw=d.textbbox((0,0),title,font=F(25))[2]; d.text(((x0+x1)/2-tw/2,y0+8),title,font=F(25),fill=(255,255,255))
        draw_center_text(d,(x0+8,y0+58,x1-8,y1-10),body,F(19))
        if i<n-1:
            ax1=x1+gap-3; ay=top+bh/2
            d.line([x1+3,ay,ax1-8,ay],fill=(120,120,120),width=4)
            d.polygon([(ax1-8,ay-7),(ax1,ay),(ax1-8,ay+7)],fill=(120,120,120))
    by=top+bh+30
    d.rounded_rectangle([margin,by,W-margin,by+60],radius=14,fill=(250,250,250),outline=(180,180,180),width=2)
    tag="贯穿主线：言必有据 · 了解反例 · 辨析命题结构   →   发展「推理能力」"
    tw=d.textbbox((0,0),tag,font=F(22))[2]; d.text((W/2-tw/2,by+16),tag,font=F(22),fill=(30,30,30))
    path=os.path.join(IMG,"图1_审辩路径_v3.png"); img.save(path,dpi=(200,200)); return path

# ---------- 图2：拐点(猪蹄)模型 ----------
def make_fig2():
    W,H=1100,620; img=Image.new("RGB",(W,H),"white"); d=ImageDraw.Draw(img); f=F(30); fs=F(24)
    yA,yC=120,500; xL,xR=120,980
    d.line([xL,yA,xR,yA],fill=(20,20,20),width=4)
    d.line([xL,yC,xR,yC],fill=(20,20,20),width=4)
    d.text((xL-70,yA-18),"A",font=f,fill=(20,20,20)); d.text((xR+12,yA-18),"B",font=f,fill=(20,20,20))
    d.text((xL-70,yC-18),"C",font=f,fill=(20,20,20)); d.text((xR+12,yC-18),"D",font=f,fill=(20,20,20))
    Ax,Ay=360,yA; Cx,Cy=430,yC; Ex,Ey=560,310
    d.line([Ax,Ay,Ex,Ey],fill=(200,40,40),width=4)
    d.line([Cx,Cy,Ex,Ey],fill=(200,40,40),width=4)
    for xx in range(140,int(Ex),22):
        d.line([xx,Ey,min(xx+11,Ex),Ey],fill=(30,120,200),width=3)
    d.text((150,Ey-44),"F",font=f,fill=(30,120,200))
    d.text((150,Ey+8),"（辅助线 EF ∥ AB）",font=fs,fill=(30,120,200))
    d.ellipse([Ex-7,Ey-7,Ex+7,Ey+7],fill=(200,40,40)); d.text((Ex+14,Ey-8),"E",font=f,fill=(200,40,40))
    d.ellipse([Ax-6,Ay-6,Ax+6,Ay+6],fill=(20,20,20)); d.ellipse([Cx-6,Cy-6,Cx+6,Cy+6],fill=(20,20,20))
    d.text((Ax-6,Ay+12),"∠BAE",font=fs,fill=(120,90,0))
    d.text((Cx-20,Cy-46),"∠DCE",font=fs,fill=(120,90,0))
    d.text((Ex+16,Ey+20),"∠AEC",font=fs,fill=(150,20,20))
    path=os.path.join(IMG,"图2_拐点模型.png"); img.save(path,dpi=(200,200)); return path

FIG1=make_fig1(); FIG2=make_fig2()
print("图1:",FIG1); print("图2:",FIG2)


# ================= 文档 =================
doc = Document()
sec = doc.sections[0]
sec.page_width=Cm(21.0); sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.5); sec.bottom_margin=Cm(2.5); sec.left_margin=Cm(2.7); sec.right_margin=Cm(2.7)
normal=doc.styles['Normal']; normal.font.name=SONG; normal.font.size=Pt(10.5)
normal._element.rPr.rFonts.set(qn('w:eastAsia'),SONG)
LINE=1.5

pt=doc.add_paragraph(); pt.alignment=WD_ALIGN_PARAGRAPH.CENTER; pt.paragraph_format.space_after=Pt(6)
r=pt.add_run("让学生审辩AI的证明"); set_cn_font(r,HEI,17,bold=True)
psub=doc.add_paragraph(); psub.alignment=WD_ALIGN_PARAGRAPH.CENTER; psub.paragraph_format.space_after=Pt(10)
r=psub.add_run("——依托“三个助手”、指向推理能力的“命题与证明”习题讲评课"); set_cn_font(r,HEI,12.5,bold=True)
pa=doc.add_paragraph(); pa.alignment=WD_ALIGN_PARAGRAPH.CENTER; pa.paragraph_format.space_after=Pt(2)
r=pa.add_run("〔作者姓名〕"); set_cn_font(r,KAI,11)
paf=doc.add_paragraph(); paf.alignment=WD_ALIGN_PARAGRAPH.CENTER; paf.paragraph_format.space_after=Pt(10)
r=paf.add_run("（上海外国语大学云间中学，上海　201600）"); set_cn_font(r,KAI,9)

def label_para(doc,label,segs,size=9.5):
    p=doc.add_paragraph(); p.paragraph_format.line_spacing=1.4; p.paragraph_format.space_after=Pt(2)
    r=p.add_run(label); set_cn_font(r,HEI,size,bold=True)
    for text,style in segs:
        rr=p.add_run(text); set_cn_font(rr,KAI,size)
        if style=='sup': rr.font.superscript=True; rr.font.name=SONG
    return p

# 摘要（≤250字）
label_para(doc,"摘要：",[
    ("针对生成式AI代答诱发的认知卸载",''),("[4][5]",'sup'),
    ("，本文反转师生与AI的关系：在“命题与证明”习题讲评课上让学生审辩AI的证明。教师借“三个助手”"
     "备课端让AI真实生成含漏洞的证明，遴选为共享错证；课中全班沿“问‘凭什么’”路径讲评辨析，"
     "教学助手同步统计每人判断；课后学生独立作证，经作业辅导助手接受AI实时挑刺与互评。"
     "共享错证保证可靠聚焦，实时活AI供给各异挑战，并与平台全员留痕互补，使“以教为学”",''),("[1][8]",'sup'),
    ("规模化落地，指向推理能力。本文属教学设计与课例研究，给出评价量规与预期，真实数据待补，不夸大。",''),
])
label_para(doc,"关键词：",[("命题与证明；几何证明；以教为学；推理能力；生成式人工智能",'')])
print("标题摘要就绪")

# ---------- 1 问题提出 ----------
heading(doc,"1　问题的提出")
heading(doc,"1.1　AI“代劳”之困：越用越不会想",size=11,before=6,after=3)
p=add_para(doc,line=LINE); runs(p,[
    ("学生把证明题拍给生成式AI，几秒即得一份“条理清晰”的解答，看似高效，却把本应自己承担的思考"
     "“卸载”给了机器，题目稍作变形便无从下手。已有研究指出，对生成式AI的过度依赖会挤占必要的认知投入，"
     "形成“认知卸载”乃至“认知债务”",''),("[4][5]",'sup'),
    ("。对以“讲道理、明因果”为命脉的几何证明而言尤其致命——它抽掉了证明教学最该留给学生的那部分挣扎。"
     "几何证明历来是初中推理训练的核心，也是难点，学生常表现出分析能力弱、推理策略少、迁移意识不足",''),
    ("[3]",'sup'),("，若再被AI代劳，推理的锻炼便无从谈起。",''),
])
heading(doc,"1.2　一个反转：从“AI替我想”到“我审AI的证明”",size=11,before=6,after=3)
p=add_para(doc,line=LINE); runs(p,[
    ("问题不在AI，而在关系。若把AI从“替学生解题的家教”改设为“证明常出漏洞、需要学生审辩的对象”，"
     "学生便从答案的接收者，变成讲证明、判真假、追依据的人。据此，本文设计一条“问‘凭什么’”审辩路径："
     "让全班审辩一个会真实犯错的AI所写的证明，把每一步都追问到“有据”。需强调，本设计不把AI的错误"
     "写死成剧本，而是利用它真实、不稳定的推理短板——它常给出“看着对、某步却站不住”的证明。"
     "为兼顾可靠与真实，核心辨析材料取自教师课前遴选的AI真实错证（全班共享、可复现），"
     "而AI“不可预料的实时回应”则放在学生独立作证后的“AI挑刺”环节与课堂追加追问中。"
     "学生要在真假难辨中把那一步找出来、讲清楚，这正是发展推理能力的关键契机。",''),
])
print("1节就绪")


# ---------- 2 立意与依据 ----------
heading(doc,"2　立意与依据")
heading(doc,"2.1　课标之锚：把落点定在“推理能力”",size=11,before=6,after=3)
p=add_para(doc,line=LINE); runs(p,[
    ("《义务教育数学课程标准（2022年版）》把“推理能力”列为初中核心素养主要表现之一，其内涵为"
     "“能够从事实和命题出发，依据规则推出结论；能够通过特殊事例推断一般结论；了解命题的结构与联系，"
     "探索并表述论证过程”",''),("[7]",'sup'),
    ("。本课任务——揪出并补全证明缺失的依据、用反例判断命题真假、辨析原命题与逆命题、独立完成论证——"
     "逐条对应上述内涵，落点集中而不泛化：主要发展“会用数学的思维思考现实世界”中的推理能力。",''),
])
heading(doc,"2.2　教参之据：言必有据、了解反例、纠错型习题",size=11,before=6,after=3)
p=add_para(doc,line=LINE); runs(p,[
    ("沪教版七年级“命题与证明”的教学要求强调：证明要“言必有据”，学生要了解“举反例可判断一个命题为假”，"
     "并能区分原命题与逆命题；教参在教学建议中提倡使用“纠错型”习题，让学生在辨析错误中体会严谨"
     "（要点归纳，精确表述以教参与课标原文为准）。审辩AI证明，正是一种“活的、动态生成的”纠错型学习。",''),
])
heading(doc,"2.3　学理之基：以教为学与“门徒效应”",size=11,before=6,after=3)
p=add_para(doc,line=LINE); runs(p,[
    ("“为教而学”能加深理解：当人预期要把知识讲给别人并真的去讲时，会更主动地组织知识、监控自身漏洞，"
     "此即“门徒效应”（protégé effect）",''),("[8]",'sup'),
    ("。让学生“讲题”的实践也表明，讲的过程能把思维暗处照亮",''),("[1]",'sup'),
    ("。本课学生对AI并非旁观式的“批改”：追问依据、补全证明、纠正假命题、再独立作证接受挑刺，"
     "本质上都是以审辩与纠正的方式“教”。几何证明本就是发展推理能力的关键载体，"
     "其教学尤重逻辑关系的显性化与问题引导",''),("[2][3]",'sup'),
    ("；把AI设为需要学生讲清、纠正的对象，恰好把这种“讲理”高频化。",''),
])
heading(doc,"2.4　两个“非它不可”：AI的活错误与平台的全员留痕",size=11,before=6,after=3)
p=add_para(doc,line=LINE); runs(p,[
    ("有三个质疑须正面回应。其一，“用一叠纸质错题卡不也一样?”——区别在AI错误的“活”：纸质错解由教师"
     "生造、往往刻意，学生易背套路；而AI的错误是其真实推理的产物，藏得深、被追问还会辩解，更像真实的思维失误，"
     "逼学生真辨。本设计不依赖AI算得对，恰恰用其逻辑不稳的短板，AI辅导的相关设计亦表明其价值更多在于"
     "对话与追问而非给出正确答案",''),("[6]",'sup'),
    ("。这里须说清“共享”与“各异”的分工，以免自相矛盾：作为核心的环节三，有意采用教师课前遴选的"
     "“共享错证”，全班辨析同一个高质量样本，以保证可靠、聚焦、可复现；而AI“各异、不可预料的实时回应”，"
     "则体现在环节五（学生独立证交AI挑刺，每人所得反馈不同）与课堂的追加追问中——二者分工不同，并不冲突，"
     "且遴选样本的“错”仍由AI真实产生，绝非教师杜撰。其二，“为何讲评AI的错，而不是学生自己的错?”——"
     "讲评学生错解固然有针对性，但AI错证可课前设计、精准命中“辅助线为什么平行”“内错角认哪个”等共性易错点，"
     "可复现、可控，又不点名暴露某位学生，且把学生从“被讲评”转为主动“审辩”。其三，“不就是平台推送吗、做的人很多?”"
     "——关键不在推送，而在“审辩活错误 + 全员留痕”：依托上海市中小学“三个助手”，教学助手把每名学生“哪一步没依据”"
     "的判断即时统计，让原本只有少数学生显性的说理过程“人人被看见”。AI供给“活的错误”，平台供给“全员的留痕”，"
     "二者互补，才使规模化的“门徒效应”真正落地。",''),
])
print("2节就绪")


# ---------- 3 教学实施 ----------
heading(doc,"3　教学实施")
heading(doc,"3.1　课型定位与组织形式",size=11,before=6,after=3)
p=add_para(doc,line=LINE); runs(p,[
    ("本课为1课时，课型是习题课中的讲评（纠错辨析）类，即习题讲评课",''),("[1]",'sup'),
    ("。它区别于以传授新知为主的新授课：以典型问题为载体，功能在于暴露思维、纠正错误、规范表述、"
     "深化理解与发展能力，而非机械重复练习。本课的创新之处，是把“讲评”的对象由传统的“学生错解”"
     "改为“AI生成的错证”。课时安排在学生已学完“命题与证明”基本规范与平行线的判定和性质之后——"
     "此定位规避“初学者悖论”：学生已具备“证明需依据、假命题可用反例否定、命题有原逆之分”的基础，"
     "才有能力去审、去辨。组织上采用“教师投屏实时审辩 + 三个助手全程支撑”：全班共审投屏的AI证明，"
     "每名学生借平板在教学助手上同步作答、标注，教师据统计即时调控。",''),
])
heading(doc,"3.2　“三个助手”的分工与技术说明",size=11,before=6,after=3)
p=add_para(doc,line=LINE); runs(p,[
    ("备课助手（课前）：以“刚学证明、爱跳步、被追问先辩解”的引导语让AI多次作答拐点模型题，"
     "从其真实输出中遴选含隐蔽漏洞的证明作为课堂素材——错误确由AI产生，又因预选而可复现"
     "（引导语见附录）。教学助手（课中）：推送审辩任务与互动题，收集并统计每名学生的真假判断、"
     "漏洞标注，支持白板注释与讨论。作业辅导助手（课后）：接收学生独立证明、组织互评与个性化反馈。"
     "投屏对话宜选不外显长思维链的通用模型，以免抢走学生思考；教师端操作、不采集学生个人数据，合规可控。",''),
])
heading(doc,"3.3　五个环节（见图1）",size=11,before=6,after=3)
pic=doc.add_paragraph(); pic.alignment=WD_ALIGN_PARAGRAPH.CENTER; pic.paragraph_format.space_after=Pt(2)
pic.add_run().add_picture(FIG1,width=Cm(15.8))
cap=doc.add_paragraph(); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER; cap.paragraph_format.space_after=Pt(6)
r=cap.add_run("图1　“问‘凭什么’”审辩路径：全班实时共审 + 平台个体留痕"); set_cn_font(r,KAI,9)

def env(doc,name,segs):
    p=add_para(doc,indent_chars=2,line=LINE,space_after=2)
    r=p.add_run(name); set_cn_font(r,HEI,10.5,bold=True)
    for text,style in segs:
        rr=p.add_run(text); set_cn_font(rr,SONG,10.5)
        if style=='i': rr.font.italic=True; rr.font.name='Times New Roman'

env(doc,"环节一　立徒弟。",[
    ("投屏出示AI并与学生约定：“今天它是刚学证明的同学，我们全班当审稿人；目标不是问倒它，"
     "而是把它证明的每一步都追问到‘有据’。”明确评价导向：只看谁把道理讲得清、依据追得准。",''),
])
env(doc,"环节二　判真假、举反例。",[
    ("投屏让AI判断假命题“若∠1 = ∠2，则∠1与∠2是对顶角”。此类简单真假题AI多能判为假——"
     "说明本课不赌“AI一定犯错”。每名学生在教学助手上作答“真/假并给出反例”，平台即时统计；"
     "再由学生画出规范反例（两个都等于40°却不对顶的角），讲清“一个反例即可否定一个命题”。",''),
])
env(doc,"环节三　揪漏洞、补依据（核心）。",[
    ("投屏出示AI对拐点（“猪蹄”）模型的证明（图2，来自备课助手遴选的含漏洞版）：已知AB∥CD，"
     "点E在两平行线之间，求证 ∠AEC = ∠BAE + ∠DCE。AI作辅助线EF∥AB，却",''),
])
p=doc.paragraphs[-1]
runs(p,[
    ("在两处留“看着对”的漏洞：一是直接用“EF∥CD”而不说明凭什么（缺“平行于同一直线的两条直线互相平行”"
     "这一依据）；二是把∠BAE与∠AEF、∠DCE与∠CEF的“内错角”关系一带而过甚至指认错位。"
     "全班逐处逼问“这一步凭什么”，每名学生同时在教学助手/白板上标出“他觉得站不住的那一步”并提交，"
     "平台统计漏洞识别的分布；随后集体补出：由AB∥CD且EF∥AB得EF∥CD，再由两组内错角相等，"
     "得∠AEC = ∠AEF + ∠CEF = ∠BAE + ∠DCE。此环节承载本课主要的推理深度。",''),
])
pic2=doc.add_paragraph(); pic2.alignment=WD_ALIGN_PARAGRAPH.CENTER; pic2.paragraph_format.space_before=Pt(2); pic2.paragraph_format.space_after=Pt(2)
pic2.add_run().add_picture(FIG2,width=Cm(8.2))
cap2=doc.add_paragraph(); cap2.alignment=WD_ALIGN_PARAGRAPH.CENTER; cap2.paragraph_format.space_after=Pt(6)
r=cap2.add_run("图2　拐点（“猪蹄”）模型及辅助线EF∥AB"); set_cn_font(r,KAI,9)

env(doc,"环节四　辨原命题与逆命题。",[
    ("投屏让AI写出“对顶角相等”的逆命题“相等的角是对顶角”并判真假。无论AI判对判错，任务都指向“讲清为什么”："
     "学生用反例（两个相等但不对顶的角）辨明——原命题与逆命题不必同真，体会命题的结构与联系；"
     "个体判断经教学助手作答留痕。",''),
])
env(doc,"环节五　独立证、平台提交。",[
    ("角色再反转：学生独立完成一道同类变式（如E在两线同侧）的证明，经作业辅导助手提交，同伴互评（点赞、评论）。"
     "教师择典型投屏，可再交AI“像老师一样挑刺”，学生复核完善。最后集体提炼“每步要有据、假命题用反例、"
     "逆命题要单独判断”等规范，回扣“问‘凭什么’”这一路径。",''),
])
p=add_para(doc,line=LINE,space_after=2); runs(p,[
    ("（教学实录截图见图3、图4：图3为教学助手对“漏洞识别”作答的实时统计（呼应全员留痕）；"
     "图4为学生独立证明在作业辅导助手上的提交与互评。二图为课堂真实截图，待正式实施后补入。）",''),
])
print("3节就绪")


# ---------- 4 评价设计与预期 ----------
heading(doc,"4　评价设计与预期表现")
p=add_para(doc,line=LINE); runs(p,[
    ("本课以过程性评价为主，且借“三个助手”把每名学生的表现留痕、可统计。笔者所带班级47人"
     "（期末90分以上22人，整体中上）。下表把教学目标、评价证据、平台工具与权重对齐，"
     "体现“教—学—评”一致；真实分层数据待实施后据实采集，不作虚报。",''),
])
heading(doc,"4.1　“说理有据”评价量规：维度×证据来源×权重",size=11,before=6,after=3)
tbl=doc.add_table(rows=5,cols=3); tbl.alignment=WD_TABLE_ALIGNMENT.CENTER; tbl.style='Table Grid'
data1=[["评价维度（对应推理能力）","权重","证据来源（“三个助手”工具）"],
       ["揪漏洞·追依据（依规则推理）","40%","环节三漏洞标注、追问发言 → 教学助手统计、白板"],
       ["用反例判真假（特殊到一般）","20%","环节二、四互动作答 → 教学助手统计"],
       ["独立证明的规范性（表述论证）","30%","环节五独立证明 → 作业辅导助手提交"],
       ["参与讲理（交流表达）","10%","讨论发言、同伴互评 → 讨论区、作业辅导助手"]]
for i,row in enumerate(data1):
    for j,val in enumerate(row):
        cell=tbl.cell(i,j); cell.text=""
        pp=cell.paragraphs[0]; pp.alignment=WD_ALIGN_PARAGRAPH.CENTER if j!=2 else WD_ALIGN_PARAGRAPH.LEFT
        rr=pp.add_run(val); set_cn_font(rr,HEI if i==0 else SONG,9,bold=(i==0))
capt=doc.add_paragraph(); capt.alignment=WD_ALIGN_PARAGRAPH.CENTER; capt.paragraph_format.space_before=Pt(2); capt.paragraph_format.space_after=Pt(6)
r=capt.add_run("表1　“说理有据”过程性评价量规（维度×证据来源×权重，权重合计100%）"); set_cn_font(r,KAI,9)

heading(doc,"4.2　核心维度的水平分档（评分参照）",size=11,before=4,after=3)
tbl2=doc.add_table(rows=4,cols=2); tbl2.alignment=WD_TABLE_ALIGNMENT.CENTER; tbl2.style='Table Grid'
data2=[["水平","可观察的行为表现"],
       ["水平三（有据）","能为每一步指明定义/公理/定理依据；主动用反例否定假命题；能辨析原、逆命题并独立完成证明"],
       ["水平二（半据）","多数步骤能说理，个别关键步仍以“显然”带过；反例与逆命题意识需教师提示"],
       ["水平一（待强化）","倾向照搬AI答案或只看结论对错；被追问依据时说理断裂"]]
for i,row in enumerate(data2):
    for j,val in enumerate(row):
        cell=tbl2.cell(i,j); cell.text=""
        pp=cell.paragraphs[0]; pp.alignment=WD_ALIGN_PARAGRAPH.CENTER if j==0 else WD_ALIGN_PARAGRAPH.LEFT
        rr=pp.add_run(val); set_cn_font(rr,HEI if i==0 else SONG,9,bold=(i==0))
capt2=doc.add_paragraph(); capt2.alignment=WD_ALIGN_PARAGRAPH.CENTER; capt2.paragraph_format.space_before=Pt(2); capt2.paragraph_format.space_after=Pt(6)
r=capt2.add_run("表2　核心维度“说理有据”水平分档（评价工具，非成效数据）"); set_cn_font(r,KAI,9)

heading(doc,"4.3　预期的可观察变化与效度边界",size=11,before=4,after=3)
p=add_para(doc,line=LINE); runs(p,[
    ("依机制预期三类变化，均需用上表与平台统计验证，不预设结论：其一，“凭什么”成为追问习惯，"
     "并从审AI迁移到互评与自评；其二，面对命题先想“能否举反例推翻”；其三，逆命题不再被想当然。"
     "须诚实说明边界：本设计尚无前后测与对照，单节课难以断言推理能力的稳定提升；平台统计虽提供了"
     "全员个体的过程数据，但其意义仍需结合作品分析审慎解读。合理定位是：它提供了一个让学生"
     "“持续处在讲理、辨错位置”的高频情境，素养发展仍需常态化积累检验。",''),
])

# ---------- 5 反思与边界 ----------
heading(doc,"5　反思与边界")
heading(doc,"5.1　AI与平台，各自不可替代",size=11,before=6,after=3)
p=add_para(doc,line=LINE); runs(p,[
    ("本设计的价值来自二者互补：AI提供“活的、可辩驳的错误”——即时生成、各不相同、会辩解，"
     "逼学生真辨而非套模板；平台提供“全员的留痕”——让每名学生的判断被看见、可统计，"
     "把只有少数人显性的说理过程还给全班。缺一则弱：只有AI而无平台，难有全员个体证据；"
     "只有平台而无AI的活错误，就退回常见的错题推送。",''),
])
heading(doc,"5.2　边界与风险",size=11,before=4,after=3)
p=add_para(doc,line=LINE); runs(p,[
    ("其一，AI可能“过度顺从”被一句话说服乱认错，或临场直接给出严谨证明——前者靠引导语“据理才认错”，"
     "后者靠“备课助手课前遴选真实错例”兜底，核心材料可复现、不赌现场。其二，须防课堂滑向“逗AI玩”，"
     "评价始终锁定“讲得有没有据”。其三，AI生成内容须教师把关，防止错误依据被当作正确知识固化。"
     "其四，需落实平板、网络与账号，投屏对话在教师端进行、不采集学生个人数据。其五，本设计定位为习题讲评课，"
     "适用于已完成新授、学生具备基本证明规范的阶段，不宜照搬到新授课。",''),
])
heading(doc,"5.3　可迁移性",size=11,before=4,after=3)
p=add_para(doc,line=LINE); runs(p,[
    ("“全班实时审辩一个会真实犯错的AI + 平台全员留痕”这一“问‘凭什么’”结构，可迁移到其他重说理与规范的内容，"
     "如代数恒等变形的依据、不等式性质的使用条件等。共性是把AI的“不严谨”转化为学生必须去辨、去讲清的对象，"
     "让“以教为学”在人人参与、过程可见中持续发生。",''),
])
print("4-5节就绪")


# ---------- 参考文献 ----------
heading(doc,"参考文献",size=12,before=10,after=4)
refs=[
    "[1]　肖巧荣.学生讲题中的教与学——以“实际问题与二元一次方程组”的习题讲评课为例[J].中学数学教学参考,2023(15):17-18.",
    "[2]　吕小勇.初中数学几何证明教学中培养学生逻辑思维的策略探究——以“四边形内角和证明”为例[J].教师,2025(【期号待补】):74-76.",
    "[3]　石鸿宾.启悟教学中的初中数学思维培养研究——以“几何证明”为例[J].数学教学通讯,2025(【期号待补】):63,86.",
    "[4]　孙发勤,李娜娜,沈霞娟,等.认知增强抑或认知债务:生成式人工智能影响认知的双重机制[J].电化教育研究,2026(9):26-33.",
    "[5]　王帆,王诗琪,史芹芹.破解大学生生成式人工智能过度依赖的原理与策略——适度认知卸载区构建与实现路径[J].电化教育研究,2026,47(8):71-79.",
    "[6]　武法提,夏志文,闫洪全.基于波利亚问题解决理论的数学辅导智能体的设计与开发研究[J].中小学数字化教学,2026(7):40-44.",
    "[7]　中华人民共和国教育部.义务教育数学课程标准(2022年版)[M].北京:北京师范大学出版社,2022:【推理能力所在页码待补】.",
    "[8]　CHASE C C, CHIN D B, OPPEZZO M A, et al. Teachable agents and the protégé effect: increasing the effort towards learning[J]. Journal of Science Education and Technology, 2009, 18(4): 334-352.",
]
for ref in refs:
    p=doc.add_paragraph(); p.paragraph_format.line_spacing=1.35; p.paragraph_format.space_after=Pt(1)
    p.paragraph_format.left_indent=Pt(21); p.paragraph_format.first_line_indent=Pt(-21)
    r=p.add_run(ref); set_cn_font(r,SONG,9)
p=add_para(doc,indent_chars=0,line=1.3,space_after=0); p.paragraph_format.space_before=Pt(3)
r=p.add_run("　注：[2][3]期号请按知网卷期补齐（页码已核）；[7]请补“推理能力”所在页码。上海市“三个助手”为平台名，正文说明即可，可不单列文献。"); set_cn_font(r,KAI,8.5,color=RGBColor(0x88,0x88,0x88))

# ---------- 附录 ----------
heading(doc,"附录　用于“备课助手”生成错例的引导语",size=12,before=10,after=4)
lines=[
    "（用途：教师在备课助手中粘贴此段，让AI以“初学者”口吻多次作答目标证明题，再遴选其中含隐蔽漏洞者作课堂素材；非学生端自定义智能体。）",
    "请你扮演一名刚开始学习几何证明、成绩中等的初中生，来完成下面这道证明题。要求：",
    "1. 爱跳步：写证明时可用“显然”“很明显”一笔带过，作辅助线后常不说明它为什么成立；",
    "2. 常想当然：容易认为“原命题成立，逆命题也一定成立”，也容易把“看起来对”的命题当真命题；",
    "3. 回应简短，像一个真实的、会犯错的同学，不要长篇展示完整推理过程；",
    "4. 若被追问“这一步凭什么”，先辩解、再想办法，只有对方给出正确反例或指明你缺了明确依据（定义/公理/定理）时才认错并改正。",
]
for ln in lines:
    p=add_para(doc,indent_chars=0,line=1.4,space_after=1); r=p.add_run(ln); set_cn_font(r,KAI,9.5)
p=add_para(doc,indent_chars=0,line=1.4,space_after=0); p.paragraph_format.space_before=Pt(8)
r=p.add_run("　说明：本稿标注“示例/待补/待核实”处，为待课堂真实实施后据实补充的数据、截图与文献信息，正式投稿前请逐项核实，不作夸大。")
set_cn_font(r,KAI,8.5,color=RGBColor(0x88,0x88,0x88))

doc.save(OUT)
print("已保存:",OUT)
