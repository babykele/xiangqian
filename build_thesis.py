# -*- coding: utf-8 -*-
"""论文版：让AI的错误成为数学课的资源 —— 输出 让AI的错误成为数学课的资源_论文.docx"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BASE="/Users/chengqp/Desktop/tan"; IMG=os.path.join(BASE,"图")
OUT=os.path.join(BASE,"让AI的错误成为数学课的资源_论文.docx")
FIG1=os.path.join(IMG,"图1_镶嵌项目路径.png")
SONG,HEI,KAI="宋体","黑体","楷体"; LINE=1.5

def setf(r,cn=SONG,size=10.5,bold=False,color=None):
    r.font.name=cn; r._element.rPr.rFonts.set(qn('w:eastAsia'),cn); r.font.size=Pt(size); r.font.bold=bold
    if color: r.font.color.rgb=color
def para(doc,ind=2,after=2,line=LINE):
    p=doc.add_paragraph(); pf=p.paragraph_format
    if ind: pf.first_line_indent=Pt(10.5*ind)
    pf.space_after=Pt(after); pf.space_before=Pt(0)
    pf.line_spacing=line; pf.line_spacing_rule=WD_LINE_SPACING.MULTIPLE
    return p
def runs(p,segs,size=10.5):
    for t,st in segs:
        r=p.add_run(t); setf(r,SONG,size)
        if st=='b': r.font.bold=True
        elif st=='sup': r.font.superscript=True
    return p
def H(doc,t,size=12,before=9,after=4):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after)
    r=p.add_run(t); setf(r,HEI,size,bold=True); return p
def h2(doc,t):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(3)
    r=p.add_run(t); setf(r,HEI,11,bold=True); return p
def dlg(doc,who,text,ai=False):
    p=para(doc,ind=1,after=1)
    r=p.add_run(who+"："); setf(r,HEI,10,bold=True,color=(RGBColor(0x1F,0x6F,0xB2) if ai else None))
    r=p.add_run(text); setf(r,SONG,10)

doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21.0); sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.5); sec.bottom_margin=Cm(2.5); sec.left_margin=Cm(2.7); sec.right_margin=Cm(2.7)
nm=doc.styles['Normal']; nm.font.name=SONG; nm.font.size=Pt(10.5); nm._element.rPr.rFonts.set(qn('w:eastAsia'),SONG)

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(5)
r=p.add_run("让AI的错误成为数学课的资源"); setf(r,HEI,17,bold=True)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(9)
r=p.add_run("——以七年级“综合与实践”《从传统连续纹样到现代镶嵌图案》为例"); setf(r,HEI,12,bold=True)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(2)
r=p.add_run("〔作者姓名〕"); setf(r,KAI,11)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(8)
r=p.add_run("（上海外国语大学云间中学，上海　201600）"); setf(r,KAI,9)

def lab(doc,label,segs,size=9.5):
    p=doc.add_paragraph(); p.paragraph_format.line_spacing=1.4; p.paragraph_format.space_after=Pt(2)
    r=p.add_run(label); setf(r,HEI,size,bold=True)
    for t,st in segs:
        rr=p.add_run(t); setf(rr,KAI,size)
        if st=='sup': rr.font.superscript=True; rr.font.name=SONG
    return p

# 摘要：问题—主张—做法—结论
lab(doc,"摘要：",[
    ("AI进数学课堂，多用来辅助讲授或生成资源，少有人把它出的错当作学习对象。本文提出一个主张："
     "AI的教学价值不在于它算得对，而在于它错得真——那些“看着合理、数学上非法”的输出，"
     "提供了传统教具给不了的审辩对象。以七年级“综合与实践”的镶嵌为例：教师课前从AI画的图中挑出“伪镶嵌”；"
     "课上让学生先猜哪些正多边形能密铺，再用“拼接点角和是否为360°”证明，进而鉴定AI的图、自己创作；"
     "学生判断经“三个助手”留痕，教师据统计当场调整教法。把AI的错变成数学判定的对象，"
     "既让几何推理有真用处，也让学生学会不轻信AI。",''),
])
lab(doc,"关键词：",[("综合与实践；镶嵌（密铺）；推理能力；生成式人工智能；批判性思维",'')])
print("头部就绪"); doc.save(OUT)


# 体裁说明
p=para(doc,ind=0,after=4); r=p.add_run("【说明】")
setf(r,HEI,9,bold=True,color=RGBColor(0xC0,0x00,0x00))
r=p.add_run("本文属教学设计与思考，第4节所呈现的学生表现为笔者预演与初步观察，未作前后测与对照；"
            "标注“示例”处待常态实施后据实补充，不夸大AI的作用。")
setf(r,KAI,9,color=RGBColor(0xC0,0x00,0x00))

# ---------- 1 问题提出：研究缺口 ----------
H(doc,"1　问题的提出")
h2(doc,"1.1　AI进了数学课，我们大多只让它“帮忙”")
para(doc); runs(doc.paragraphs[-1],[
    ("这两年AI进课堂，常见的用法有两类：帮老师备课、生成练习和课件；或者当学生的答疑助手。"
     "共同点是——都把AI当成“做得对”的帮手。可只要用得多，就会发现它常常不对：让它画一幅镶嵌图案，"
     "出来的图远看严丝合缝，凑近一瞧却有缝隙、或者叠在一起；有时它还一本正经地说“这是用正五边形密铺的”，"
     "而正五边形怎么拼也铺不满。道理不难：文生图这类AI是照着看过的海量图片把画面“画”像，"
     "并不按几何规则去“算”（会写代码、会调几何工具的AI另说）。",''),
])
para(doc); runs(doc.paragraphs[-1],[
    ("平时我们把这类错当麻烦，赶紧绕开。但对数学课来说，这里恰好有一块没人用过的资源："
     "学生要是只看图、又信了AI那份笃定，就很容易把“看着对”当成“真的对”",''),("[4]",'sup'),
    ("；而要判断它到底对不对，非得动用数学不可。",''),
])
h2(doc,"1.2　一个还没被认真用起来的空子")
para(doc); runs(doc.paragraphs[-1],[
    ("夏雪梅等人比较中美的数学项目后发现，我们的数学项目常把情境“净化”得太干净，"
     "也不太看重批判性思维",''),("[3]",'sup'),
    ("。这话说到了痛处：真实世界里的问题多半带着杂音、甚至带着错误，而我们给学生的情境往往过于干净、"
     "答案唯一。既然AI的错是真实发生、且学生往后天天会碰到的，为什么不把它请进课堂，当成要学生动数学去查验的对象?",''),
])
para(doc); runs(doc.paragraphs[-1],[
    ("由此，本文提出一个主张：",''),
    ("生成式AI在数学课上的教学价值，不在于它算得对，而在于它错得真——那些“视觉上合理、数学上非法”的输出，"
     "提供了传统教具给不了的审辩对象。",'b'),
    ("下面以七年级“综合与实践”的镶嵌内容为例，说明这个主张怎么落到一节课里，以及它为什么值得一试。",''),
])

# ---------- 2 主张的依据 ----------
H(doc,"2　这个主张站得住吗")
h2(doc,"2.1　先自问一句：这节课非用AI不可吗")
para(doc); runs(doc.paragraphs[-1],[
    ("镶嵌本是沪教版七年级“综合与实践”里的老内容，讲的是图形的平移、旋转、轴对称",''),("[2]",'sup'),
    ("。备课时笔者先问了自己一句：非用AI不可吗?说实话，不用AI，拿几张图片、剪几张纸照样能上"
     "——教材配的空中课堂就是这么上的。要是AI只拿来找找图、生成几个图案，它就是个可有可无的点缀。",''),
])
para(doc); runs(doc.paragraphs[-1],[
    ("但按前面的主张改一下重心，AI就非有不可了：这节课要审的，就是它自己画出来的东西。"
     "这跟老师故意画张错图不同——老师画错，学生一眼看出是“考我的”；AI是真不懂、还画得理直气壮，"
     "而且往后学生天天都会碰到这样的AI。也就是说，抽掉AI，这件事本身就不存在了。",''),
])
h2(doc,"2.2　为什么镶嵌是个好载体")
para(doc); runs(doc.paragraphs[-1],[
    ("因为它的对错没有余地。图案好不好看可以各有所爱，但能不能铺满平面，只看一条：拼接点处各角加起来"
     "是不是恰好360°，且不重叠。这是一把刻度分明的尺子，最适合用来戳穿“看着对、其实错”。"
     "对七年级学生来说，这尺子又刚好够得着——只要会算多边形的内角。",''),
])
h2(doc,"2.3　这节课到底在练什么")
para(doc); runs(doc.paragraphs[-1],[
    ("这节课归在“综合与实践”，用项目的方式上",''),("[1]",'sup'),
    ("，但练的还是数学：算内角、用360°判断能不能密铺，练的是推理；对着AI的图挑毛病、发现它站不住，"
     "练的是敢质疑、会质疑——这两样课标里都写着。至于“学会不轻信AI”，是学生在这节数学课里顺带得到的一份见识，"
     "不必硬说成什么数学素养。这节课的正事是数学，AI只是把学生引到数学跟前的那个由头。",''),
])
print("1-2节就绪"); doc.save(OUT)


# ---------- 3 实践：三个论证小节 ----------
H(doc,"3　怎么做：让AI的错变成可用的资源")
if os.path.exists(FIG1):
    pic=doc.add_paragraph(); pic.alignment=WD_ALIGN_PARAGRAPH.CENTER; pic.paragraph_format.space_after=Pt(2)
    pic.add_run().add_picture(FIG1,width=Cm(15.6))
    cp=doc.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_after=Pt(6)
    r=cp.add_run("图1　本项目的推进路线：以AI的“伪镶嵌”为对象，以数学判定为主线"); setf(r,KAI,9)
para(doc); runs(doc.paragraphs[-1],[
    ("整个项目约6课时（课内4、课外2），顺着教材的活动走：赏纹样、辨变换、还原埃舍尔的鸽子瓷砖，"
     "再到先猜后证、鉴定AI的作品、自己造一幅能密铺的图（图1）。赏纹样、辨变换、做作品这几步，"
     "用图片和剪纸也做得来；真正离不开AI的只有“鉴定”那一步。下面只讲支撑本文主张的三件事。",''),
])

h2(doc,"3.1　第一件事：让AI“真错”，而且错得可控")
para(doc); runs(doc.paragraphs[-1],[
    ("既然要审AI的错，就得先让它错，还不能指望它上课时恰好犯错。笔者的做法是课前“策展”："
     "用备课助手连着下指令——“画一幅用正五边形密铺的图案”“生成一幅四方连续的镶嵌”“把正六边形变形成动物再密铺”，"
     "一口气出十来张；再按三条标准挑出3张留着上课用：一是缝隙或重叠肉眼看得见（放大后更明显），"
     "二是错在正多边形的误用（尤其正五边形），三是整体“看着挺像”，不能一眼就假。",''),
])
para(doc); runs(doc.paragraphs[-1],[
    ("这样处理有个好处：错是AI自己犯的，不是老师编的；同时因为事先挑过，课堂上不必赌运气，"
     "别的班再上一次也能复现。要提醒的是，如果直接用对话模型去问“正五边形能不能密铺”，"
     "现在多数模型会答“不能”——所以这一步要用的是它",''),("画出来的图",'b'),("，而不是它说出来的话。",''),
])

h2(doc,"3.2　第二件事：把“凭什么”算清楚，从查表走到证明")
para(doc); runs(doc.paragraphs[-1],[
    ("要审AI，学生手里得先有尺子。笔者没有直接给结论，而是先让全班猜：哪些正多边形能单独密铺?"
     "猜完在平台上提交。分歧最大的是正五边形和正六边形——这正是好的探究起点。接着才一起把标准立起来："
     "无缝隙、不重叠，且每个拼接点处各内角之和恰为360°。",''),
])
t=doc.add_table(rows=6,cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
d1=[["正多边形","每个内角","能否单独密铺","判定（拼接点角和=360°）"],
    ["正三角形","60°","能","60°×6 = 360°"],
    ["正方形","90°","能","90°×4 = 360°"],
    ["正五边形","108°","不能","108°×3=324°，×4=432°，无整数倍等于360°"],
    ["正六边形","120°","能","120°×3 = 360°"],
    ["正八边形（拓展）","135°","不能单独","与正方形组合：135°×2+90°=360°"]]
for i,row in enumerate(d1):
    for j,v in enumerate(row):
        c=t.cell(i,j); c.text=""; pp=c.paragraphs[0]
        pp.alignment=WD_ALIGN_PARAGRAPH.CENTER if j!=3 else WD_ALIGN_PARAGRAPH.LEFT
        rr=pp.add_run(v); setf(rr,HEI if i==0 else SONG,9,bold=(i==0))
cp=doc.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_before=Pt(2); cp.paragraph_format.space_after=Pt(6)
r=cp.add_run("表1　正多边形能否单独密铺（末行为拓展）"); setf(r,KAI,9)
para(doc); runs(doc.paragraphs[-1],[
    ("表1只是结果，不能停在这里。这节课真正的数学分量，在于带学生把“为什么偏偏是正三、正四、正六”算明白："
     "正n边形每个内角是(n−2)·180°/n，要能单独密铺，每个拼接点就得由整数个这样的角拼满360°，"
     "即 360°÷[(n−2)·180°/n] = 2n/(n−2) 必须是正整数。逐一看：n=3、4、6时，这个值分别是6、4、3，都是整数；"
     "n=5时是10/3，不是整数——这正是正五边形铺不满的根子；n≥7时，2n/(n−2) 夹在2与3之间"
     "（因为 2n/(n−2)=2+4/(n−2)，而 0<4/(n−2)≤4/5<1），也不可能取到整数。所以",''),
    ("能单独密铺的正多边形只有正三角形、正方形、正六边形三种。",'b'),
    ("从“查表”走到“算明白”，学生回头审AI的图才真正有底气。",''),
])
p=para(doc,ind=2,after=4); r=p.add_run("需要提醒："); setf(r,HEI,9.5,bold=True,color=RGBColor(0xC0,0x00,0x00))
r=p.add_run("上述判定只对“正”多边形。“正五边形不能密铺”不等于“五边形不能密铺”——有多类不规则凸五边形是可以密铺的。"
            "若学生查到并提出，应当肯定，并借此说清“正多边形”与“一般多边形”的分野，作为课后拓展。")
setf(r,KAI,9.5,color=RGBColor(0xC0,0x00,0x00))

h2(doc,"3.3　第三件事：让学生从“看着像”走到“算得准”")
para(doc); runs(doc.paragraphs[-1],[
    ("尺子有了，就把课前挑好的AI作品投上去。下面这一张，是笔者让AI按“画一幅用正五边形密铺的图案”"
     "这条指令画出来的。",''),
])
dlg(doc,"教师","这是AI照“用正五边形密铺”画出来的图。它做到了吗?别急着看，先拿数学量一量。")
dlg(doc,"学生","正五边形一个内角是 (5−2)×180°÷5 = 108°。")
dlg(doc,"学生","拼接点得凑满360°。可是108×3=324不到，108×4=432又超了，没有整数个刚好360!")
dlg(doc,"教师","所以这张图会怎么样?")
dlg(doc,"学生","拼接的地方一定留缝，铺不满。")
dlg(doc,"教师","放大看看（图2）。每个拼接点果然都有一道缝——它画得像，可画的是假密铺。")
dlg(doc,"学生","换正六边形就行，120×3=360，能铺。")
p=para(doc,ind=0,after=3); r=p.add_run("〔图2位置〕")
setf(r,HEI,9,bold=True,color=RGBColor(0xC0,0x00,0x00))
r=p.add_run("AI按“用正五边形密铺”指令生成的图案及其拼接点放大（可见缝隙）——此图由教师用备课助手出图后截取，待补入。")
setf(r,KAI,9,color=RGBColor(0xC0,0x00,0x00))
para(doc); runs(doc.paragraphs[-1],[
    ("这段对话里，学生做的事已经不是“认图”，而是拿一个数量条件去否定一个看似成立的结论——"
     "这正是推理的样子。此后再看AI的图，他们的第一句话开始从“挺好看”变成“先算算它拼不拼得满”。",''),
])
print("3节就绪"); doc.save(OUT)


# ---------- 4 评价与观察 ----------
H(doc,"4　评价怎么做，看到了什么")
h2(doc,"4.1　评什么：把证据落在成果上")
para(doc); runs(doc.paragraphs[-1],[
    ("这类课最怕评成“谁的图好看”。笔者把评价挂在两样成果上：一份个人的《鉴定报告》"
     "（判AI的图合不合法、依据是什么、怎么改），一份小组的镶嵌作品加《设计说明》"
     "（单位图案怎么来、用了什么变换、拼接点角和为什么是360°）。权重与证据来源见表2。",''),
])
t=doc.add_table(rows=6,cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
d2=[["评价维度","权重","证据来源"],
    ["密铺合法性的判定（推理）","30%","《鉴定报告》中的角度计算与结论"],
    ["审辩AI：指出错处并提出修改","30%","《鉴定报告》、课堂共审发言、平台作答"],
    ["识别单位纹样与变换","15%","辨变换环节的平台作答"],
    ["作品创作与设计说明","15%","小组作品与《设计说明》"],
    ["协作与互评","10%","讨论区发言、同伴互评"]]
for i,row in enumerate(d2):
    for j,v in enumerate(row):
        c=t.cell(i,j); c.text=""; pp=c.paragraphs[0]
        pp.alignment=WD_ALIGN_PARAGRAPH.CENTER if j!=2 else WD_ALIGN_PARAGRAPH.LEFT
        rr=pp.add_run(v); setf(rr,HEI if i==0 else SONG,9,bold=(i==0))
cp=doc.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_before=Pt(2); cp.paragraph_format.space_after=Pt(6)
r=cp.add_run("表2　评价维度、权重与证据来源（合计100%）"); setf(r,KAI,9)

h2(doc,"4.2　平台不只是留个记录，它让笔者当场改了教法")
para(doc); runs(doc.paragraphs[-1],[
    ("这一点想单独说。用“三个助手”，好处不在于把学生的答案存下来，而在于笔者能当场看见"
     "——猜想那一步的统计一出来，就知道有相当一部分学生认为正五边形能密铺（示例：47人中约22人）。"
     "换在从前，举手表决只看得见几个积极的孩子，这个分歧未必浮出水面。看见了，笔者就临时加了一轮小组辩论，"
     "让持两种意见的学生各自说理由，再一起去算——教学顺序是被数据改掉的。",''),
])
para(doc); runs(doc.paragraphs[-1],[
    ("笔者以为，这才是数字平台在数学课上真正的用处：",''),
    ("把原本看不见的分歧变成看得见的分歧，让教师据此决定下一步讲什么。",'b'),
    ("否则它只是个答题器。",''),
])
h2(doc,"4.3　初步看到的变化（示例，待常态实施后补实）")
para(doc); runs(doc.paragraphs[-1],[
    ("笔者所带班级47人，整体中上。预演与初步观察中有三点变化值得记录，但都还需更多证据："
     "一是判断的起点变了，看到AI的图先算角、后下结论；二是说理有了依据，从“看着有缝”变成"
     "“108°×3只有324°，凑不满”；三是做自己的作品时会主动验一验——有一组原想用五角星当单位图案，"
     "算过发现凑不满360°就换掉了，还把这段过程写进了《设计说明》。这一条笔者尤其看重："
     "它说明学生已经把“好看”让给了“合法”。",''),
])
para(doc); runs(doc.paragraphs[-1],[
    ("为便于同行照着观察自己的班，笔者把《鉴定报告》的表现分成三层：",''),
    ("能算出内角、判定合不合法、并说出该怎么改，为第三层；能判对合不合法，但依据说不全、"
     "还停在“看着有缝”，为第二层；仍凭观感下结论、被追问依据就说不下去，为第一层。",'b'),
    ("这三层也可用于课后对作品《设计说明》的评阅。真实的分布待常态实施后据实统计。",''),
])
h2(doc,"4.4　把话说在前头")
para(doc); runs(doc.paragraphs[-1],[
    ("本文没做前后测、也没有对照班，单凭一个项目，不好说学生的推理能力就稳稳提高了。"
     "平台能给出每个人的过程数据，可这些数字究竟说明什么，还得连着他们交的《鉴定报告》一起看。"
     "稳妥的说法是：它给学生造了个反复“拿数学去查AI”的真实机会；至于素养，还得靠平时一节一节地攒。",''),
])

# ---------- 5 结论与讨论 ----------
H(doc,"5　结论与讨论")
para(doc); runs(doc.paragraphs[-1],[
    ("回到本文的主张：AI在数学课上的用处，不必总盯着它算得对不对；它错得真，反而给了数学一个"
     "非常合适的用场——让学生拿着数学的尺子去查验一个看似可信的结论。镶嵌只是个方便的入口，"
     "因为它的判据干净利落（角和是否360°）。同样的做法还能搬到别处，条件是那件事AI真的常做错、"
     "而数学又判得清：比如让它画一张尺规作图的步骤图，学生查每一步是否只用了圆规和直尺；"
     "让它按给定数据生成一张统计图，学生核对扇形的圆心角、条形的高度跟数据合不合；"
     "让它画一个所谓“黄金分割”的矩形分割，学生量一量比例对不对。这些地方AI都爱想当然，"
     "而学生手里恰好有能判它的工具。",''),
])
para(doc); runs(doc.paragraphs[-1],[
    ("也有几条得守住。第一，主角始终是数学，别让“审AI”抢了戏——学生练的是几何与推理，"
     "“学会不轻信AI”是捎带的收获。第二，判对错的权得攥在数学和学生手里，别请AI来当裁判；"
     "已有研究也说，AI在数学里的长处是陪着聊、给点启发，而不是提供标准答案",''),("[5]",'sup'),
    ("。第三，上课用的AI错例一定要课前挑过，既留着看得出的毛病，又别把学生带偏。"
     "第四，这个设计放在学生已有基本经验之后更合适；新授阶段学生尚无判断力，容易被误导。"
     "下一步，笔者将在常态教学中继续实施，把平台记录、学生的《鉴定报告》与作品作为证据积累起来。",''),
])

# ---------- 参考文献 ----------
H(doc,"参考文献")
refs=[
    "[1]　中华人民共和国教育部.义务教育数学课程标准(2022年版)[M].北京:北京师范大学出版社,2022.",
    "[2]　义务教育教科书（五·四学制）数学·七年级上册及其教学参考资料[M].上海:上海教育出版社.（综合与实践·从传统连续纹样到现代镶嵌图案）",
    "[3]　夏雪梅,刘潇.素养视角下中美数学项目驱动性问题设计的比较研究[J].全球教育展望,2022,51(7):45-61.",
    "[4]　王帆,王诗琪,史芹芹.破解大学生生成式人工智能过度依赖的原理与策略——适度认知卸载区构建与实现路径[J].电化教育研究,2026,47(8):71-79.",
    "[5]　武法提,夏志文,闫洪全.基于波利亚问题解决理论的数学辅导智能体的设计与开发研究[J].中小学数字化教学,2026(7):40-44.",
]
for ref in refs:
    p=doc.add_paragraph(); p.paragraph_format.line_spacing=1.35; p.paragraph_format.space_after=Pt(1)
    p.paragraph_format.left_indent=Pt(21); p.paragraph_format.first_line_indent=Pt(-21)
    r=p.add_run(ref); setf(r,SONG,9)
p=para(doc,ind=0,after=0,line=1.3); p.paragraph_format.space_before=Pt(6)
r=p.add_run("　注：5篇文献均在正文标注引用（[1]§2.3、[2]§2.1、[3]§1.2、[4]§1.1、[5]§5）。"
            "图2（AI生成的“伪镶嵌”及拼接点放大）请用备课助手出图后截取补入；第4节标“示例”处，"
            "待常态实施后以平台记录、学生《鉴定报告》与作品据实补充；投稿送审版请将作者与单位信息移至首页单列。")
setf(r,KAI,8.5,color=RGBColor(0x88,0x88,0x88))
doc.save(OUT); print("已保存:",OUT)
