# -*- coding: utf-8 -*-
"""生成投稿前自查清单 Word：投稿自查清单_命题与证明.docx"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BASE="/Users/chengqp/Desktop/tan"
OUT=os.path.join(BASE,"投稿自查清单_命题与证明.docx")
SONG,HEI,KAI="宋体","黑体","楷体"

def setf(run,cn=SONG,size=10.5,bold=False,color=None):
    run.font.name=cn; run._element.rPr.rFonts.set(qn('w:eastAsia'),cn)
    run.font.size=Pt(size); run.font.bold=bold
    if color: run.font.color.rgb=color

doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21.0); sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.2); sec.bottom_margin=Cm(2.2); sec.left_margin=Cm(2.5); sec.right_margin=Cm(2.5)
nm=doc.styles['Normal']; nm.font.name=SONG; nm.font.size=Pt(10.5)
nm._element.rPr.rFonts.set(qn('w:eastAsia'),SONG)

def title(t,size=15):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(4)
    r=p.add_run(t); setf(r,HEI,size,bold=True)
def sub(t):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(10)
    r=p.add_run(t); setf(r,KAI,10)
def h(t,size=12,before=10,after=4):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after)
    r=p.add_run(t); setf(r,HEI,size,bold=True)
def item(t,box="□",bold_head=None,color=None,size=10.5):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.3
    p.paragraph_format.left_indent=Pt(14); p.paragraph_format.first_line_indent=Pt(-14)
    r=p.add_run(box+" "); setf(r,SONG,size)
    if bold_head:
        r=p.add_run(bold_head); setf(r,HEI,size,bold=True,color=color)
    r=p.add_run(t); setf(r,SONG,size,color=color)
def note(t):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.3
    p.paragraph_format.left_indent=Pt(28)
    r=p.add_run(t); setf(r,KAI,9.5,color=RGBColor(0x66,0x66,0x66))

title("《让学生审辩AI的证明》投稿前自查清单")
sub("对标《上海中学数学》《中学数学教学参考》体例 ·  逐项打钩，全绿即可投")

print("头部就绪")
doc.save(OUT)

doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21.0); sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.2); sec.bottom_margin=Cm(2.2); sec.left_margin=Cm(2.5); sec.right_margin=Cm(2.5)
nm=doc.styles['Normal']; nm.font.name=SONG; nm.font.size=Pt(10.5)
nm._element.rPr.rFonts.set(qn('w:eastAsia'),SONG)
title("《让学生审辩AI的证明》投稿前自查清单")
sub("对标《上海中学数学》《中学数学教学参考》体例 · 逐项打钩，全绿即可投")

RED=RGBColor(0xC0,0x00,0x00)
h("一、必须补齐的“待补”项（只能你来定，不可杜撰）")
item("补 参考文献[2] 吕小勇 的期号（页码74-76已核）——知网条目直接抄。", bold_head="〔硬〕", color=RED)
item("补 参考文献[3] 石鸿宾 的期号与出版年（页码63,86已核；年份以版权页为准）。", bold_head="〔硬〕", color=RED)
item("补 参考文献[7] 课标中“推理能力”所在页码（翻2022版课标“核心素养内涵”一节）。", bold_head="〔硬〕", color=RED)
item("核对[1] 肖巧荣是“2023(15)”还是“2023年第5期(下旬)”——以知网导出的著录为准。", bold_head="〔硬〕", color=RED)
item("图3：教学助手对“漏洞识别”作答的实时统计截图（上课当堂采，类比井维兰案例的答题分布图）。", bold_head="〔硬〕", color=RED)
item("图4：学生独立证明在作业辅导助手上的提交与互评截图（当堂采）。", bold_head="〔硬〕", color=RED)
item("表1/表2 的真实分层数据 + §4 学生原话：上课后据实填，禁止编造，宁可写“约”与区间。", bold_head="〔硬〕", color=RED)
note("提示：以上标红项不补齐不要投；补齐后把正文里的【期号待补】【页码待补】【待补入】字样全部删净。")

h("二、上课采数据 · 现场要留的证据（照此采，回来正好填§4）")
item("环节二、四：每人“真/假 + 反例”的互动作答 → 导出统计（正确率、分布）。")
item("环节三（核心）：每人标注“哪一步没依据”的作答 → 导出“漏洞识别分布”，作图3。")
item("环节五：收齐学生独立证明（拍照/提交）→ 按表2三档评分，得真实分层人数。")
item("挑2-3句学生原话（如“为了教它，我得把‘显然’讲出来”）→ 填§4.3，标注真实来源。")
item("留课堂实录视频/关键对话截图各1份，备匿名送审与后续佐证。")
note("你班级已知：47人、期末90分以上22人——分层人数据实报，与量规三档对应。")

h("三、体例与格式（对标刊物，逐条核）")
item("标题＝主标题＋“——…以…”副标题，点明课型（习题讲评课）与素养（推理能力）。已符 ✅")
item("摘要 150~250字、第三人称、无“本文分N部分”。已控制在250字 ✅（改动后再数一次）")
item("关键词 3~5个、分号隔开、含课标术语（推理能力/命题与证明）。已符 ✅")
item("全文分节编号统一（1/2/3、1.1…），不混用。已符 ✅")
item("图1、图2、表1、表2 均有编号+题注，正文有“见图1/如表1”引出。已符 ✅")
item("参考文献 GB/T 7714，正文上标[1]…[8]，已引课标。补全待补项后即符。")
item("篇幅约3000~5000字（正文约4000字）。已符 ✅")
print("一二三节就绪")
doc.save(OUT)

h("四、内容与逻辑（评审最看的“真问题—真实践—真反思”）")
item("成效不夸大：无前后测/对照要写明；用“预期/待实证”，不作因果断言。已符 ✅")
item("“共享错证 vs 实时活AI”的分工在摘要/§1.2/§2.4/§5口径一致，不自相矛盾。已焊死 ✅")
item("已正面回应三问：与错题卡之别、为何讲评AI的错、平台是否平庸。已符 ✅")
item("课型定义规范（习题讲评课，含暴露思维/纠错/深化功能），有文献[1]锚。已符 ✅")
item("数学正确性：拐点模型 ∠AEC=∠BAE+∠DCE 的证明与依据链无误。已核 ✅")
item("AI赋能点明确落在“备课助手生成真实错证”，未写平台做不到的“学生端自定义智能体”。已符 ✅")

h("五、匿名送审与投稿事务")
item("送审版：正文正文内不得出现真实姓名、学校可识别信息；作者/单位单列首页或另页。", bold_head="〔硬〕", color=RED)
note("现稿标题下有“上外云间中学”，送审前移到单独的作者信息页。")
item("中图分类号/文献标识码/文章编号：可留空或占位，由编辑部分配。")
item("确认未一稿多投；图片清晰可复用（≥300dpi）；作者单位全称准确。")
item("择刊：教学实践类优先《上海中学数学》或《中学数学教学参考》（下旬）。")

h("六、可选加分（有余力再做）")
item("环节五加“分层选做”，让数学深度再厚一点。")
item("§4 增一张“学生独立证明”前后对照的作品图（同一人或典型样本）。")
item("正文一处点明本设计与培训所倡“审辩式、慢用AI、把思维还给学生”理念的呼应。")

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(12)
r=p.add_run("　一句话定稿标准：标红〔硬〕项全部补齐、正文再无“待补/示例/待核实”字样、成效表述有边界不夸大——达成即可投。")
setf(r,KAI,9.5,color=RGBColor(0x88,0x88,0x88))

doc.save(OUT)
print("已保存:",OUT)
