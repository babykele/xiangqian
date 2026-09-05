# -*- coding: utf-8 -*-
"""生成项目式教学案例投稿稿：镶嵌数智项目_教学案例投稿稿.docx（实施/成效为模拟预演，标注待替换）"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

BASE="/Users/chengqp/Desktop/tan"; IMG=os.path.join(BASE,"图")
OUT=os.path.join(BASE,"镶嵌数智项目_教学案例投稿稿.docx")
FIG1=os.path.join(IMG,"图1_镶嵌项目路径.png")
SONG,HEI,KAI="宋体","黑体","楷体"; LINE=1.5

def setf(r,cn=SONG,size=10.5,bold=False,color=None):
    r.font.name=cn; r._element.rPr.rFonts.set(qn('w:eastAsia'),cn); r.font.size=Pt(size); r.font.bold=bold
    if color: r.font.color.rgb=color
def para(doc,ind=2,after=2,line=LINE,align=None):
    p=doc.add_paragraph(); pf=p.paragraph_format
    if ind: pf.first_line_indent=Pt(10.5*ind)
    pf.space_after=Pt(after); pf.space_before=Pt(0); pf.line_spacing=line; pf.line_spacing_rule=WD_LINE_SPACING.MULTIPLE
    if align: p.alignment=align
    return p
def runs(p,segs,size=10.5):
    for t,st in segs:
        r=p.add_run(t); setf(r,SONG,size)
        if st=='b': r.font.bold=True
        elif st=='sup': r.font.superscript=True
    return p
def H(doc,t,size=12,before=8,after=4):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after)
    r=p.add_run(t); setf(r,HEI,size,bold=True); return p
def h2(doc,t):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(3)
    r=p.add_run(t); setf(r,HEI,11,bold=True); return p

doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21.0); sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.5); sec.bottom_margin=Cm(2.5); sec.left_margin=Cm(2.7); sec.right_margin=Cm(2.7)
nm=doc.styles['Normal']; nm.font.name=SONG; nm.font.size=Pt(10.5); nm._element.rPr.rFonts.set(qn('w:eastAsia'),SONG)

# 标题
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(5)
r=p.add_run("当AI画错了镶嵌"); setf(r,HEI,17,bold=True)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(9)
r=p.add_run("——用几何“密铺律”识破生成式AI的“伪镶嵌”：一则“综合与实践”项目式教学案例"); setf(r,HEI,12,bold=True)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(2)
r=p.add_run("〔作者姓名〕"); setf(r,KAI,11)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(8)
r=p.add_run("（上海外国语大学云间中学，上海　201600）"); setf(r,KAI,9)

def lab(doc,label,segs,size=9.5):
    p=doc.add_paragraph(); p.paragraph_format.line_spacing=1.4; p.paragraph_format.space_after=Pt(2)
    r=p.add_run(label); setf(r,HEI,size,bold=True)
    for t,st in segs:
        rr=p.add_run(t); setf(rr,KAI,size)
        if st=='sup': rr.font.superscript=True; rr.font.name=SONG
    return p

lab(doc,"摘要：",[
    ("让AI画一幅镶嵌图案，它常画得很像、其实拼不满，甚至用正五边形“密铺”——可它根本铺不满。"
     "本案例抓住这种“看着对其实错”的AI作品，依托沪教版七年级“综合与实践”"
     "《从传统连续纹样到现代镶嵌图案》，把重心从“用AI学镶嵌”换成“用数学检验AI画得对不对”："
     "学生赏纹样、辨变换，猜哪些正多边形能密铺、用“角和是否360°”验证，再鉴定AI的图，"
     "最后自己造一幅能密铺的图；分组、作答、互评都在“三个助手”平台完成。学生由此明白："
     "好不好看是一回事，能不能铺满得靠数学算。（实施与数据为模拟预演，待真实上课替换。）",''),
])
lab(doc,"关键词：",[("综合与实践；镶嵌（密铺）；图形的运动；生成式人工智能；创新意识",'')])
print("头部就绪"); doc.save(OUT)


# 醒目声明
p=para(doc,ind=0,after=4); r=p.add_run("【体裁说明】")
setf(r,HEI,9,bold=True,color=RGBColor(0xC0,0x00,0x00))
r=p.add_run("本稿为项目式教学案例的写作稿；第3节“实施过程”与第4节“成效”为设计阶段的模拟预演"
            "（对话与统计均为模拟），正式实施后须以真实课堂实录、平台数据与学生作品替换，不夸大AI作用。")
setf(r,KAI,9,color=RGBColor(0xC0,0x00,0x00))

# ---------- 1 问题提出 ----------
H(doc,"1　问题的提出")
h2(doc,"1.1　AI“会画不会算”：以假乱真的“伪镶嵌”")
para(doc); runs(doc.paragraphs[-1],[
    ("现在的AI画图很唬人。让它“画一幅镶嵌图案”，出来的图远看严丝合缝，凑近一瞧却常有缝隙、或者叠在一起；"
     "有时它还一本正经地说“这是用正五边形密铺的”，可正五边形怎么拼也铺不满。道理不难：文生图这类AI是照着"
     "看过的海量图片把画面“画”像，并不会按几何规则去“算”（会写代码、会调几何工具的AI另说）。"
     "学生要是只看图、又信了AI那份笃定，就很容易把“看着对”当成“真的对”。这种对AI的盲信，正是笔者担心的",''),("[4]",'sup'),("。",''),
])
h2(doc,"1.2　从“用AI学镶嵌”到“用数学识破AI”：一次立意反转")
para(doc); runs(doc.paragraphs[-1],[
    ("镶嵌本是沪教版七年级“综合与实践”里的老内容，讲的是图形的平移、旋转、轴对称",''),("[2]",'sup'),
    ("。备课时笔者先问了自己一句：这节课非用AI不可吗?说实话，不用AI，拿几张图片、剪几张纸照样能上"
     "——教材配的空中课堂就是这么上的。要是AI只拿来找找图、生成几个图案，那它就是个可有可无的点缀，谈不上什么新意。",''),
])
para(doc); runs(doc.paragraphs[-1],[
    ("于是笔者把想法倒了过来：",''),
    ("这节课不是“用AI帮我们学镶嵌”，而是“拿数学去检验AI画的镶嵌到底对不对”。",'b'),
    ("这样一来，AI就非有不可了——因为整节课要审的，就是它自己画出来的东西。这跟老师故意画张错图不一样："
     "老师画错，学生一眼看出是“考我的”；AI是真不懂、还画得理直气壮，而且往后学生天天都会碰到这样的AI。"
     "用什么来判断?几何给了一把最准的尺子：能不能密铺，就看拼接点的角加起来够不够360°，一点都含糊不得。",''),
])
para(doc); runs(doc.paragraphs[-1],[
    ("这么改，也正好补了一个大家都认的短板。夏雪梅等人比较中美的数学项目后发现，我们的数学项目"
     "常把情境“净化”得太干净，也不太看重批判性思维",''),("[3]",'sup'),
    ("。而“审AI画的镶嵌”这件事，情境是真的（AI真会这么错），要的又正是批判地去质疑——这两点恰恰是平时缺的。",''),
])
print("1节就绪"); doc.save(OUT)


# ---------- 2 项目设计 ----------
H(doc,"2　项目设计")
h2(doc,"2.1　课标与素养定位（厘清两个层次）")
para(doc); runs(doc.paragraphs[-1],[
    ("这节课归在“综合与实践”，用项目的方式来上",''),("[1]",'sup'),
    ("。真正练的还是数学：算内角、拿360°判断能不能密铺，练的是推理；对着AI的图挑毛病、发现它站不住，"
     "练的是敢质疑、会质疑——这两样课标里都写着（推理能力、创新意识）。至于“学会不轻信AI”，"
     "是学生在这节数学课里顺带得到的一份见识，不必硬说成什么数学素养。一句话：这节课的正事是数学，"
     "AI只是把学生引到数学跟前的那个由头。",''),
])
h2(doc,"2.2　驱动问题")
para(doc); runs(doc.paragraphs[-1],[
    ("“AI画的镶嵌，能信吗?”——面对AI生成、看上去天衣无缝的“镶嵌”图，能不能用数学找出破绽、说清它错在哪、"
     "再造出真正合法的镶嵌?",'b'),
])
para(doc); runs(doc.paragraphs[-1],[
    ("拿夏雪梅那套框架",''),("[3]",'sup'),
    ("回头看这个问题：它是真事（AI确实会这么画，不是为讲课编出来哄学生的），跟技术和生活都沾边，"
     "想答清楚得又算又说、不算浅，而且逼着学生去质疑——正好补上前面说的那两块短板。",''),
])
h2(doc,"2.3　数学内核：镶嵌（密铺）的判定标准")
para(doc); runs(doc.paragraphs[-1],[
    ("镶嵌合法的两条数学标准：①无缝隙、不重叠地铺满平面；②每个拼接点处各内角之和恰为360°。"
     "由此判定正多边形能否单独密铺（表1），这正是学生审辩AI作品的“数学尺子”。",''),
])
para(doc); runs(doc.paragraphs[-1],[
    ("不必满足于查表，还可以带学生把“为什么偏偏是正三、正四、正六”算清楚：正n边形每个内角是(n−2)·180°/n，"
     "要能单独密铺，就得让它整除360°，也就是 360°÷[(n−2)·180°/n] = 2n/(n−2) 是正整数。"
     "n取3、4、6时正好得6、4、3；n到7以后，这个值不到3也不是整数了。所以",''),
    ("能单独密铺的正多边形，只有正三角形、正方形、正六边形三种。",'b'),
    ("从“查表”走到“算明白”，这节课的数学分量就上来了，学生回头审AI的图也才真正有底气。",''),
])
t=doc.add_table(rows=6,cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
d1=[["正多边形","每个内角","能否单独密铺","数学判定（拼接点角和=360°）"],
    ["正三角形","60°","能","60°×6 = 360°"],
    ["正方形","90°","能","90°×4 = 360°"],
    ["正五边形","108°","不能","108°×3=324°，×4=432°，无整数倍等于360°"],
    ["正六边形","120°","能","120°×3 = 360°"],
    ["正八边形（拓展）","135°","不能单独","与正方形组合：135°×2+90°=360°（半正密铺）"]]
for i,row in enumerate(d1):
    for j,v in enumerate(row):
        c=t.cell(i,j); c.text=""; pp=c.paragraphs[0]
        pp.alignment=WD_ALIGN_PARAGRAPH.CENTER if j!=3 else WD_ALIGN_PARAGRAPH.LEFT
        rr=pp.add_run(v); setf(rr,HEI if i==0 else SONG,9,bold=(i==0))
cp=doc.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_before=Pt(2); cp.paragraph_format.space_after=Pt(6)
r=cp.add_run("表1　正多边形能否单独密铺的数学判定（审辩AI作品的依据；末行为拓展）"); setf(r,KAI,9)
p=para(doc,ind=2,after=4); r=p.add_run("教师须知："); setf(r,HEI,9.5,bold=True,color=RGBColor(0xC0,0x00,0x00))
r=p.add_run("本课判定严格针对“正”多边形。须注意“正五边形不能密铺”不等于“五边形不能密铺”——"
            "存在多类不规则凸五边形可以密铺平面；若学生提出此点，应予肯定，并指明这正是“正多边形”与“一般多边形”的分野，可作课后拓展。")
setf(r,KAI,9.5,color=RGBColor(0xC0,0x00,0x00))

h2(doc,"2.4　六环节路径与“三个助手”分工")
if os.path.exists(FIG1):
    pic=doc.add_paragraph(); pic.alignment=WD_ALIGN_PARAGRAPH.CENTER; pic.paragraph_format.space_after=Pt(2)
    pic.add_run().add_picture(FIG1,width=Cm(15.8))
    cp=doc.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_after=Pt(6)
    r=cp.add_run("图1　镶嵌审辩课六环节路径：AI“伪密铺”为情境、数学判定为主线"); setf(r,KAI,9)
para(doc); runs(doc.paragraphs[-1],[
    ("整个项目走六步：①赏纹样（越王勾践剑纹、唐草纹）②辨变换（认平移、旋转、轴对称）③还原埃舍尔的鸽子瓷砖"
     "④先猜后证（猜哪些正多边形能密铺，再用360°验）⑤给AI的作品做鉴定（★核心，这一步离不开AI）"
     "⑥自己造一幅能密铺的图。一共约6课时（课内4、课外2），顺着教材的活动走，重头在第⑤步。"
     "其实①②⑥用图片、剪纸也做得来，真正非用AI不可的，只有第⑤步——要审的，得是AI亲手画的东西。",''),
])
t=doc.add_table(rows=4,cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
d2=[["阶段","助手","主要任务"],
    ["课前","备课助手","生成传统纹样素材；用引导语让AI生成“待审的镶嵌图”（含缺陷）；备分层任务单"],
    ["课中","教学助手","推“识别变换”“猜想”互动并实时统计；白板标注；投屏共审AI作品，收集每人判断"],
    ["课后","作业辅导助手","接收小组镶嵌作品与设计说明；组织同伴互评与个性化反馈"]]
for i,row in enumerate(d2):
    for j,v in enumerate(row):
        c=t.cell(i,j); c.text=""; pp=c.paragraphs[0]
        pp.alignment=WD_ALIGN_PARAGRAPH.CENTER if j!=2 else WD_ALIGN_PARAGRAPH.LEFT
        rr=pp.add_run(v); setf(rr,HEI if i==0 else SONG,9,bold=(i==0))
cp=doc.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_before=Pt(2); cp.paragraph_format.space_after=Pt(6)
r=cp.add_run("表2　“三个助手”在项目各阶段的分工"); setf(r,KAI,9)
print("2节就绪"); doc.save(OUT)


def dlg(doc,who,text,ai=False):
    p=para(doc,ind=1,after=1)
    r=p.add_run(who+"："); setf(r,HEI,10,bold=True,color=(RGBColor(0x1F,0x6F,0xB2) if ai else None))
    r=p.add_run(text); setf(r,SONG,10)
def stat(doc,text):
    p=para(doc,ind=1,after=2); r=p.add_run("〔模拟平台数据〕"); setf(r,KAI,9,color=RGBColor(0x99,0x66,0x00))
    r=p.add_run(text); setf(r,KAI,9,color=RGBColor(0x99,0x66,0x00))

# ---------- 3 实施过程 ----------
H(doc,"3　实施过程（模拟预演，待真实实录替换）")
para(doc); runs(doc.paragraphs[-1],[
    ("本项目约6课时（课内4＋课外2）。下按课时呈现推进主线；对话与统计均为设计阶段的模拟，"
     "正式实施后以真实实录与平台截图替换（拟配图2“教学助手实时统计”、图3“学生镶嵌作品与互评”）。",''),
])
h2(doc,"3.1　课前准备：备课助手“策展”待审图")
para(doc); runs(doc.paragraphs[-1],[
    ("上课前，笔者用备课助手让AI多画几张——“画一幅正五边形密铺”“生成一幅四方连续的镶嵌”，一口气出十来张，"
     "再从里头挑3张“看着像、其实拼不满”的留着上课用（有的正五边形之间露着缝，有的是假对称，有的干脆错位）。"
     "另外备几张传统纹样图、一份分层任务单，和《鉴定报告》《设计说明》两张空白模板。学生四人一组、组内好中弱搭配，"
     "一人一台平板登平台。这些错图都是AI自己画的、又是笔者事先挑好的，既真实，上课也不至于翻车。",''),
])
h2(doc,"3.2　第一课时·赏纹样与辨变换：先审AI的“识别”")
para(doc); runs(doc.paragraphs[-1],[
    ("先从越王勾践剑上的菱形纹说起，看看唐草纹（二方连续）、清代的福寿纹（四方连续）；学生来了兴致，"
     "就进入“辨变换”——这回让AI先开口“认”，学生来挑它的错。",''),
])
dlg(doc,"教师","（投屏一条“回纹”二方连续）先请AI说说它的单位纹样和变换方式。")
dlg(doc,"AI","这条纹样是把一个“回”字形单位纹样不断平移排列而成的。",ai=True)
dlg(doc,"教师","同意吗?各组在教学助手上作答“它用了哪些变换”。")
stat(doc,"平移 28人／平移＋旋转(180°) 15人／含轴对称 4人——多数与AI一样只看到平移。")
dlg(doc,"学生1","不全是平移，相邻两个“回”是反着的，应该有180°旋转。")
dlg(doc,"教师","上来比一比：平移→不重合；旋转180°→重合。AI只说对了一半——对不对，得用变换的定义量。")

h2(doc,"3.3　第二课时·还原埃舍尔与提猜想·学裁判")
para(doc); runs(doc.paragraphs[-1],[
    ("这节课先接着教材的活动3：把埃舍尔那只不规则的鸽子瓷砖，通过平移、轴对称、旋转，还原成规规矩矩的基本图形，"
     "再铺开去。铺着铺着学生自己就会问：到底怎样才算真铺满、才算真密铺?笔者就势把问题抛出来——先猜，再证。",''),
])
dlg(doc,"教师","先别算，你们猜：哪些正多边形能单独密铺?在教学助手上写下猜想。")
stat(doc,"正三角形能 41人／正方形能 45人／正五边形能 22人／正六边形能 26人——正五、正六分歧最大。")
dlg(doc,"教师","怎么不靠猜、把它定下来?")
dlg(doc,"学生","看拼接点凑不凑得满360°!")
dlg(doc,"教师","对。各角之和恰好360°、且无缝不重叠——拿这把尺子去量表1里每一个。")
dlg(doc,"学生","那我猜正五边形能，是错的，得改。")

h2(doc,"3.4　第三课时·审AI作品：识破“伪镶嵌”（核心）")
dlg(doc,"教师","我让AI“画一幅用正五边形密铺的图案”，它交卷了。")
dlg(doc,"AI","（给出一张正五边形铺展图）这是用正五边形密铺的图案，它们无缝拼接。",ai=True)
dlg(doc,"教师","真无缝吗?先别信，拿数学量。正五边形一个内角多少?")
dlg(doc,"学生","(5−2)×180°÷5 = 108°。")
dlg(doc,"学生2","拼接点要360°。可108×3=324不到、108×4=432又超，没有整数个刚好360!")
dlg(doc,"教师","放大看——（投屏放大）每个拼接点果然留着一道缝。AI会画，画的却是“假密铺”。")
stat(doc,"“这幅AI图合法吗”：开始判“不合法”19人；经角度计算＋放大验证后判“不合法”44人，3人待助。")
dlg(doc,"学生3","换正六边形，120×3=360，能铺。")
dlg(doc,"教师","这就是数学当裁判——AI造图，我们判它合不合法。")

h2(doc,"3.5　课外·共创作品与成果产出")
dlg(doc,"学生组","我们让AI把正六边形“变”成一条小鱼当单位图案。")
dlg(doc,"教师","变形后还能密铺吗?凭什么?")
dlg(doc,"学生","变形只是把边“凹进去一块、凸出去同样一块”，拼接角度没变，拼接点还是三个120°=360°，照样无缝密铺。")
dlg(doc,"教师","AI负责把它变好看，你们负责证明它还合法——作品和说明传到作业辅导助手，互评。")
p=para(doc,ind=2,after=2); r=p.add_run("（严格版：对边界作平移式改造，割去的部分在相邻砖上原样补回，故仍无缝密铺；七年级点到“角度没变、能补回”的直觉即可。）")
setf(r,KAI,9,color=RGBColor(0x88,0x88,0x88))

h2(doc,"3.6　学生物化成果样例（模拟，待真实作品替换）")
def sample(doc,title,lines):
    p=para(doc,ind=1,after=1); r=p.add_run(title); setf(r,HEI,10,bold=True)
    for lab,txt in lines:
        pp=para(doc,ind=2,after=1)
        r=pp.add_run(lab); setf(r,HEI,9.5,bold=True)
        r=pp.add_run(txt); setf(r,SONG,9.5)
sample(doc,"样例1　学生《AI“伪镶嵌”鉴定报告》（节选）",[
    ("待审作品：","备课助手生成的“正五边形密铺”图。"),
    ("判定：","不合法（假密铺）。"),
    ("数学依据：","正五边形每个内角108°；一个拼接点最多拼3个，108°×3=324°<360°，必然留缝，故不可能无缝密铺。"),
    ("图上证据：","放大后每个拼接点均见明显缝隙。"),
    ("修改建议：","换用正六边形，120°×3=360°，可无缝密铺；或改用能密铺的不规则五边形（拓展）。"),
])
sample(doc,"样例2　小组镶嵌《设计说明》（节选）",[
    ("作品名：","《游不散的鱼》。"),
    ("单位图案：","由正六边形变形而来——将一条边向内凹去一小块、在对边向外凸出等量的一块，形似小鱼。"),
    ("所用变换：","平移（沿两个方向四方连续）。"),
    ("合法性验证：","变形只搬动边界、未改变拼接处的角；每个拼接点仍是3个120°=360°，且无缝隙不重叠，故合法。"),
    ("自评：","起初想用五角星做单位图案，验证时发现拼接点凑不满360°，遂放弃——这说明好看不等于能密铺。"),
])
print("3节就绪"); doc.save(OUT)


# ---------- 4 成效与分析 ----------
H(doc,"4　成效与分析（模拟预演数据，待真实核实）")
para(doc); runs(doc.paragraphs[-1],[
    ("笔者所带班级47人（期末90分以上22人，整体中上）。以下为依机制作出的预演性描述与示例数据，"
     "均需真实实施后据实采集，本文不作夸大结论。",''),
])
h2(doc,"4.1　过程性评价：维度×证据来源×权重")
t=doc.add_table(rows=6,cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
d3=[["评价维度（对应素养）","权重","证据来源（成果／助手）"],
    ["密铺合法性的数学判定（空间观念·推理）","30%","AI作品《鉴定报告》的角度计算与判定"],
    ["审辩AI：识破伪密铺、指错与修正（核心·创新意识）","30%","《鉴定报告》、课堂共审 → 教学助手统计"],
    ["识别单位纹样与变换（几何直观）","15%","环节②互动作答 → 教学助手统计"],
    ["镶嵌作品创作与设计说明（创新意识）","15%","小组作品与《设计说明》 → 作业辅导助手"],
    ["协作与互评（交流表达）","10%","讨论区发言、同伴互评"]]
for i,row in enumerate(d3):
    for j,v in enumerate(row):
        c=t.cell(i,j); c.text=""; pp=c.paragraphs[0]
        pp.alignment=WD_ALIGN_PARAGRAPH.CENTER if j!=2 else WD_ALIGN_PARAGRAPH.LEFT
        rr=pp.add_run(v); setf(rr,HEI if i==0 else SONG,9,bold=(i==0))
cp=doc.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_before=Pt(2); cp.paragraph_format.space_after=Pt(6)
r=cp.add_run("表3　镶嵌项目过程性评价量规（维度×证据来源×权重，合计100%）"); setf(r,KAI,9)
para(doc); runs(doc.paragraphs[-1],[
    ("就拿最看重的“审AI”来说，预演里能“判出这张是假密铺、还说得出为什么”的学生，从一开始的十来个"
     "增加到四十来个（这是预演的估计，真数字得上了课才知道）。这里的“四十来个”指“结论判对”的人，"
     "和下面表4按“说理到什么程度”分的层不是一回事——真实采集时这两个口径要分开记。"
     "每个人的作答记录、交上来的《鉴定报告》，平台都留得下。",''),
])
h2(doc,"4.2　物化成果的分层（示例数据，待实证）")
para(doc); runs(doc.paragraphs[-1],[
    ("按表3量规评《鉴定报告》与小组镶嵌作品，模拟预演中的分层示例见表4。",''),
])
t=doc.add_table(rows=4,cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
d4=[["水平","典型表现","本班分布（示例）"],
    ["水平三（有据）","能算内角、判合法性并给出可行修改；作品含正确的360°验证","约20人／约6组"],
    ["水平二（半据）","能判对合法性，但依据表述不全，个别只说“看着有缝”未落到角度","约19人"],
    ["水平一（待强化）","仍以“像不像”下判断，被追问依据时说理断裂","约8人"]]
for i,row in enumerate(d4):
    for j,v in enumerate(row):
        c=t.cell(i,j); c.text=""; pp=c.paragraphs[0]
        pp.alignment=WD_ALIGN_PARAGRAPH.CENTER if j!=1 else WD_ALIGN_PARAGRAPH.LEFT
        rr=pp.add_run(v); setf(rr,HEI if i==0 else SONG,9,bold=(i==0))
cp=doc.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_before=Pt(2); cp.paragraph_format.space_after=Pt(6)
r=cp.add_run("表4　成果水平分层（示例数据，待真实实施采集）"); setf(r,KAI,9)
para(doc); runs(doc.paragraphs[-1],[
    ("典型作品点评（模拟）：样例2《游不散的鱼》的价值不在好看，而在《设计说明》里“先试五角星、"
     "验证凑不满360°而放弃”的记录——它显示学生把“美观”让位于“数学合法”，正是本项目要的思维转变。",''),
])
h2(doc,"4.3　预期的可观察变化（待实证，不预设结论）")
para(doc); runs(doc.paragraphs[-1],[
    ("照这个设计，笔者估计会看到几样变化，但都得等真实数据来验，不敢先下结论：学生看到AI的图，"
     "第一反应从“挺好看”变成“先算算它拼不拼得满”；慢慢明白“看着对”不等于“真的对”，会把“无缝”落到"
     "“角加起来是不是360°”上；到自己做作品时，也会主动验一验合不合法。到底有没有、有多少，上完课再说。",''),
])
h2(doc,"4.4　效度的诚实边界")
para(doc); runs(doc.paragraphs[-1],[
    ("也得把话说在前头：这节课没做前后测、也没有对照班，单靠一个项目，不好说学生的素养就稳稳提高了。"
     "平台能给出每个人的过程数据，可这些数字究竟说明什么，还得连着他们交的《鉴定报告》一起看。"
     "稳妥的说法是：它给学生造了个反复“拿数学去查AI”的真实机会；至于素养，还得靠平时一节一节地攒。",''),
])

# ---------- 5 反思与建议 ----------
H(doc,"5　反思与建议")
para(doc); runs(doc.paragraphs[-1],[
    ("预演下来，笔者有几点想法。最要紧的一条：这节课的主角始终是数学，AI只是个由头，别让“审AI”抢了戏"
     "——学生练的是几何和推理，“学会不轻信AI”是捎带的收获。再一个，说这课“非AI不可”，指的是核心那一步："
     "要审的必须是AI真画的东西；至于赏纹样、辨变换、做作品那几步，用图片、剪纸也行，所以这课经得起"
     "“把AI拿掉还能不能上”的追问。还有，判对错的权得攥在数学和学生手里，别让AI来当裁判——已有研究也说，"
     "AI在数学里的长处是陪着聊、给点启发，而不是提供标准答案",''),("[5]",'sup'),
    ("。另外，上课用的AI错图，一定要笔者课前先看过、挑过，既留着看得出的毛病，又别把学生带偏；"
     "课上也得盯着，别变成“比谁画得好看”。最后，这稿里的对话和数据都是笔者预演时拟的，"
     "等真上了课，把平台记录、学生的《鉴定报告》和作品补进来，再拿去投。",''),
])

# ---------- 参考文献 ----------
H(doc,"参考文献")
refs=[
    "[1]　中华人民共和国教育部.义务教育数学课程标准(2022年版)[M].北京:北京师范大学出版社,2022.",
    "[2]　义务教育教科书（五·四学制）数学·七年级上册及其教学参考资料[M].上海:上海教育出版社.（综合与实践·从传统连续纹样到现代镶嵌图案）",
    "[3]　夏雪梅,刘潇.素养视角下中美数学项目驱动性问题设计的比较研究[J].全球教育展望,2022,51(7):45-61.",
    "[4]　王帆,王诗琪,史芹芹.破解大学生生成式人工智能过度依赖的原理与策略——适度认知卸载区构建与实现路径[J].电化教育研究,2026,47(8):71-79.",
    "[5]　武法提,夏志文,闫洪全.基于波利亚问题解决理论的数学辅导智能体的设计与开发研究[J].中小学数字化教学,2026(7):40-44.",
]
for ref in refs:
    p=doc.add_paragraph(); p.paragraph_format.line_spacing=1.35; p.paragraph_format.space_after=Pt(1)
    p.paragraph_format.left_indent=Pt(21); p.paragraph_format.first_line_indent=Pt(-21)
    r=p.add_run(ref); setf(r,SONG,9)
p=para(doc,ind=0,after=0,line=1.3); p.paragraph_format.space_before=Pt(6)
r=p.add_run("　说明：5篇文献均在正文标注引用（[1]§2.1、[2]§1.2、[3]§1.2及§2.2、[4]§1.1、[5]§5）。"
            "第3、4节数据与对话为模拟预演，须以真实实施的实录、平台统计与学生作品替换后方可投稿。")
setf(r,KAI,8.5,color=RGBColor(0x88,0x88,0x88))
doc.save(OUT); print("已保存:",OUT)
