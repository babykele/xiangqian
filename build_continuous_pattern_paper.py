# -*- coding: utf-8 -*-
"""生成《从图形运动到生成规则》论文初稿。"""

import os
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


BASE = Path("/Users/chengqp/Desktop/tan")
REF = BASE / "让AI的错误成为数学课的资源_上海中学数学修订稿.docx"
OUT = BASE / "从图形运动到生成规则_生成式人工智能支持初中生数学应用意识发展的教学实践_论文稿.docx"
WORK = Path("/private/tmp/ai_pattern_paper")
IMG_DIR = BASE / "图"
WORK.mkdir(parents=True, exist_ok=True)
IMG_DIR.mkdir(parents=True, exist_ok=True)

# 正文采用宋体风格，标题采用黑体风格；渲染时显式加载系统中文字体。
BODY_CN = "Songti SC"
HEAD_CN = "Hiragino Sans GB"
LATIN = "Times New Roman"
BLUE = RGBColor(31, 78, 121)
DARK = RGBColor(30, 30, 30)
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EDF4FA"
PALE_YELLOW = "FFF7D6"
GRID = "8FAADC"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def set_cell_width(cell, width_cm):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, cn=BODY_CN, size=10.5, bold=False, italic=False, color=DARK):
    run.font.name = LATIN
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), cn)
    run._element.rPr.rFonts.set(qn("w:ascii"), LATIN)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_widow_control(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:widowControl"))
    if node is None:
        node = OxmlElement("w:widowControl")
        p_pr.append(node)


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:keepNext"))
    if node is None:
        node = OxmlElement("w:keepNext")
        p_pr.append(node)


def add_body(doc, text="", indent=True, before=0, after=0, line=1.42, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if indent:
        pf.first_line_indent = Pt(21)
    if text:
        r = p.add_run(text)
        set_run_font(r)
    set_widow_control(p)
    return p


def add_runs(paragraph, parts, size=10.5, cn=BODY_CN):
    """parts: [(text, style)], style in normal/bold/italic/sup."""
    for text, style in parts:
        run = paragraph.add_run(text)
        set_run_font(run, cn=cn, size=size, bold=(style == "bold"), italic=(style == "italic"))
        if style == "sup":
            run.font.superscript = True
            run.font.size = Pt(max(7.5, size - 2))
    return paragraph


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    if level == 1:
        pf.space_before = Pt(9)
        pf.space_after = Pt(4)
        size = 12
    else:
        pf.space_before = Pt(6)
        pf.space_after = Pt(3)
        size = 11
    pf.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, cn=HEAD_CN, size=size, bold=True, color=BLUE)
    keep_with_next(p)
    return p


def add_label_para(doc, label, parts, size=9.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label)
    set_run_font(r, cn=HEAD_CN, size=size, bold=True)
    add_runs(p, parts, size=size)
    return p


def add_table(doc, headers, rows, widths, font_size=9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for j, (header, width) in enumerate(zip(headers, widths)):
        cell = table.rows[0].cells[j]
        set_cell_width(cell, width)
        set_cell_shading(cell, LIGHT_BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.08
        r = p.add_run(header)
        set_run_font(r, cn=HEAD_CN, size=font_size, bold=True, color=BLUE)
    for i, row_data in enumerate(rows):
        new_row = table.add_row()
        prevent_row_split(new_row)
        cells = new_row.cells
        for j, (value, width) in enumerate(zip(row_data, widths)):
            cell = cells[j]
            set_cell_width(cell, width)
            if i % 2 == 1:
                set_cell_shading(cell, "F8FAFC")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j > 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(value)
            set_run_font(r, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    keep_with_next(p)
    r = p.add_run(text)
    set_run_font(r, size=9, color=RGBColor(80, 80, 80))
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=9, color=RGBColor(100, 100, 100))


def arrow(draw, x1, y1, x2, y2, color=(88, 111, 137), width=5):
    draw.line((x1, y1, x2, y2), fill=color, width=width)
    import math
    angle = math.atan2(y2-y1, x2-x1)
    length = 16
    spread = 0.52
    pts = [(x2, y2),
           (x2-length*math.cos(angle-spread), y2-length*math.sin(angle-spread)),
           (x2-length*math.cos(angle+spread), y2-length*math.sin(angle+spread))]
    draw.polygon(pts, fill=color)


def make_flow_figure():
    font_path = "/System/Library/Fonts/Hiragino Sans GB.ttc"
    f_title = ImageFont.truetype(font_path, 34)
    f_box = ImageFont.truetype(font_path, 27)
    f_small = ImageFont.truetype(font_path, 22)
    image = Image.new("RGB", (1800, 720), "white")
    d = ImageDraw.Draw(image)
    title = "从现实纹样到可检验生成物：学生主导的数学应用闭环"
    tb = d.textbbox((0, 0), title, font=f_title)
    d.text(((1800-(tb[2]-tb[0]))/2, 26), title, font=f_title, fill=(31, 78, 121))

    labels = [
        ("真实需求", "明确用途与限制"),
        ("识别结构", "找基本图案与重复单元"),
        ("建立模型", "确定运动方式与要素"),
        ("生成规则", "把模型写成可执行条件"),
        ("AI执行", "输出带网格或可编辑图"),
        ("数学检验", "查对应点、接缝与不变量"),
        ("修正迁移", "改规则并解释适用边界"),
    ]
    colors = [(226, 239, 250), (231, 244, 238), (238, 236, 250), (255, 243, 218),
              (233, 241, 247), (255, 231, 231), (229, 245, 242)]
    xs = [40, 290, 540, 790, 1040, 1290, 1540]
    top, bw, bh = 150, 220, 220
    for i, ((head, sub), x) in enumerate(zip(labels, xs)):
        d.rounded_rectangle((x, top, x+bw, top+bh), radius=24, fill=colors[i], outline=(95, 126, 155), width=3)
        hb = d.textbbox((0, 0), head, font=f_box)
        d.text((x+(bw-(hb[2]-hb[0]))/2, top+38), head, font=f_box, fill=(30, 30, 30))
        line1, line2 = sub[:7], sub[7:]
        lines = [line1] + ([line2] if line2 else [])
        y = top+108
        for line in lines:
            sb = d.textbbox((0, 0), line, font=f_small)
            d.text((x+(bw-(sb[2]-sb[0]))/2, y), line, font=f_small, fill=(70, 70, 70))
            y += 35
        if i < len(xs)-1:
            arrow(d, x+bw+5, top+bh/2, xs[i+1]-8, top+bh/2)

    d.rounded_rectangle((280, 460, 1520, 620), radius=28, fill=(248, 250, 252), outline=(155, 170, 185), width=3)
    d.text((330, 492), "认知分工", font=f_box, fill=(31, 78, 121))
    d.text((520, 487), "学生：建模—表达—验证—修正", font=f_box, fill=(25, 25, 25))
    d.text((520, 545), "AI：执行—呈现—暴露歧义", font=f_box, fill=(25, 25, 25))
    d.text((1030, 516), "检验不通过", font=f_small, fill=(173, 47, 47))
    arrow(d, 1280, 525, 930, 410, color=(173, 47, 47), width=4)
    path = IMG_DIR / "连续纹样_应用意识发展路径.png"
    image.save(path, dpi=(220, 220))
    return path


FIGURE = make_flow_figure()


# 以原稿为版式来源，保留页面设置和可复用样式，正文全部重建。
work_docx = WORK / "continuous_pattern_work.docx"
shutil.copy2(REF, work_docx)
doc = Document(work_docx)
body = doc._element.body
for child in list(body):
    if child.tag != qn("w:sectPr"):
        body.remove(child)

sec = doc.sections[0]
sec.start_type = WD_SECTION_START.NEW_PAGE
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.2)
sec.bottom_margin = Cm(2.2)
sec.left_margin = Cm(2.5)
sec.right_margin = Cm(2.5)
sec.header_distance = Cm(1.1)
sec.footer_distance = Cm(1.1)

# 清理页眉页脚，重建简洁页码。
for paragraph in sec.header.paragraphs:
    paragraph._element.getparent().remove(paragraph._element)
for paragraph in sec.footer.paragraphs:
    paragraph._element.getparent().remove(paragraph._element)
add_page_number(sec.footer.add_paragraph())

normal = doc.styles["Normal"]
normal.font.name = LATIN
normal.font.size = Pt(10.5)
normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_CN)

props = doc.core_properties
props.title = "从图形运动到生成规则：生成式人工智能支持初中生数学应用意识发展的教学实践——以‘连续纹样’为例"
props.subject = "初中数学综合与实践；生成式人工智能；应用意识"
props.author = ""
props.last_modified_by = ""
props.keywords = "连续纹样,生成式人工智能,应用意识,图形运动,数学建模"
props.comments = ""


# 标题区
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
r = p.add_run("从图形运动到生成规则：")
set_run_font(r, cn=HEAD_CN, size=17, bold=True, color=RGBColor(0, 0, 0))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(7)
p.paragraph_format.line_spacing = 1.2
r = p.add_run("生成式人工智能支持初中生数学应用意识发展的教学实践")
set_run_font(r, cn=HEAD_CN, size=11.5, bold=True, color=RGBColor(0, 0, 0))
r.add_break()
r = p.add_run("——以“连续纹样”为例")
set_run_font(r, cn=HEAD_CN, size=11.5, bold=True, color=RGBColor(0, 0, 0))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(1)
r = p.add_run("〔作者姓名〕")
set_run_font(r, size=10.5)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(7)
r = p.add_run("（〔单位名称〕，上海　〔邮编〕）")
set_run_font(r, size=9.5)

add_label_para(doc, "摘要：", [(
    "连续纹样是上海初中数学“图形的运动”单元中的综合与实践内容。针对生成式人工智能容易把活动窄化为“快速出图”的问题，本文以校园边饰设计为真实任务，构建“识别基本图案—建立图形运动模型—形成五要素生成规则—AI执行—数学检验—迭代修正”的学习链。学生负责提取结构、确定变换要素、表达规则并验证结果，AI负责批量执行和显化规则歧义；评价聚焦模型合理性、规则精确性、证据充分性及迁移水平。该设计把AI由答案提供者转化为模型执行器和可质疑对象，使图形运动知识经历从现实情境到数学表达、再返回现实检验的完整过程，为应用意识发展提供可观察的学习证据。",
    "normal")], size=9.5)
add_label_para(doc, "关键词：", [("生成式人工智能；应用意识；连续纹样；图形运动；生成规则", "normal")], size=9.5)

note = doc.add_table(rows=1, cols=1)
note.alignment = WD_TABLE_ALIGNMENT.CENTER
note.autofit = False
cell = note.cell(0, 0)
set_cell_width(cell, 15.7)
set_cell_shading(cell, PALE_YELLOW)
set_cell_margins(cell, top=90, bottom=90, start=120, end=120)
p = cell.paragraphs[0]
p.paragraph_format.line_spacing = 1.15
r = p.add_run("稿件说明：本文为依据课程标准、上海教材及相关研究形成的教学实践稿。文中课堂片段为预设性演绎，不作为已经发生的课堂实录；正式投稿前须用真实学生作品、对话记录和评价数据替换或印证相应内容。")
set_run_font(r, size=8.8, color=RGBColor(90, 72, 24))


# 1 问题提出
add_heading(doc, "1　问题提出")
add_heading(doc, "1.1　从“生成图案”转向“生成规则”", level=2)
p = add_body(doc)
add_runs(p, [
    ("“连续纹样”兼具生活性、审美性与数学性。上海七年级上册“图形的运动”综合与实践要求学生从装饰图案中识别基本图案，分析平移、旋转、轴对称等形成规律，并在重构与创作中发展几何直观、空间观念和创新意识", "normal"),
    ("[1][2]", "sup"),
    ("。然而，生成式人工智能进入课堂后，最容易发生的并不是“数学更深”，而是学生输入“帮我生成一个连续纹样”，随即得到一幅看似完整的图。图案虽然丰富，图形运动的对象、方向、距离以及重复边界却被隐藏在机器结果中。学生完成了作品，却未必经历把现实问题数学化的过程。", "normal")
])
p = add_body(doc)
add_runs(p, [
    ("因此，本文不把“AI会画连续纹样”作为创新，而把教学重心前移到规则形成：学生能否从现实纹样中提取可重复结构，能否把图形运动模型表达为AI可执行、同伴可理解、结果可检验的规则；当AI的输出与意图不一致时，能否依据数学关系解释偏差并修正规则。这里的“生成”首先是学生生成数学关系，其次才是AI生成图案。", "normal")
])

add_heading(doc, "1.2　相关研究的启示与可推进空间", level=2)
p = add_body(doc)
add_runs(p, [
    ("已有平面镶嵌或纹样项目多从真实设计任务、数学与艺术融合、作品评价等角度组织学习。李佩乔等以校园设计节为载体，形成指向应用意识的项目式学习路径", "normal"),
    ("[4]", "sup"),
    ("；苏国东借助电子白板复制图形，促进学生观察拼接关系与形成规律", "normal"),
    ("[5]", "sup"),
    ("。这类研究为情境、任务和技术操作提供了成熟经验，也意味着仅以“纹样创作+技术辅助”为题，创新空间已经有限。", "normal")
])
p = add_body(doc)
add_runs(p, [
    ("生成式人工智能研究则多聚焦备课、资源生成、对话答疑或综合与实践的一般流程。赵文君等提出“需求分析—策略制定—工具匹配—指令设置—评估优化”的概念教学路径，并提醒防止认知外包", "normal"),
    ("[6]", "sup"),
    ("；沈爱桐等论述AI支持任务分析、资源供给与评价反馈的实践路径", "normal"),
    ("[7]", "sup"),
    ("。但在学生端，如何把数学模型转译为生成约束、如何用数学证据否定一幅“视觉上像对的”AI图案，仍缺少细化的课堂机制。本文据此提出研究问题：生成式人工智能如何作为“规则执行器”和“歧义显化器”，支持初中生经历数学化、模型应用与结果解释的完整过程？", "normal")
])


# 2 课程依据与理论框架
add_heading(doc, "2　课程依据与理论框架")
add_heading(doc, "2.1　应用意识：把图形运动用于解释与设计", level=2)
p = add_body(doc)
add_runs(p, [
    ("《义务教育数学课程标准（2022年版）》把应用意识列为初中阶段核心素养的主要表现，强调有意识地运用数学概念、原理和方法解释现实世界中的现象、解决现实问题，并能感悟数学与现实的联系", "normal"),
    ("[1]", "sup"),
    ("。“综合与实践”倡导在真实情境中提出问题、作出假设、建立模型、检验和修正模型，形成可交流的成果。由此，连续纹样活动的素养落点不应只是认识平移、轴对称和旋转，更要考查学生是否主动调用这些知识解释纹样形成规律，并据此完成受条件约束的设计。", "normal")
])

add_caption(doc, "表1　教学任务与课程标准相关要求的对应")
add_table(doc,
          ["学习环节", "学生的数学活动", "主要素养表现", "可观察证据"],
          [
              ("识别结构", "从实物纹样中确定基本图案和重复单元", "几何直观、空间观念", "在图中圈画单元并说明判据"),
              ("建立模型", "选择平移、旋转或轴对称，确定运动要素", "应用意识、模型观念", "用图示、文字或式子表示对应关系"),
              ("规则表达", "把模型转化为完整、无歧义的生成条件", "数学表达、推理能力", "五要素规则卡及同伴复述结果"),
              ("检验修正", "依据对应点、边界和不变量核验AI输出", "应用意识、推理能力", "标注证据、指出偏差、提出修正规则"),
              ("迁移设计", "将规则用于新图案并解释适用边界", "创新意识、应用意识", "成品、设计说明与迁移任务"),
          ],
          [2.2, 5.0, 3.0, 5.5], font_size=8.6)

add_heading(doc, "2.2　数学建模循环：不以“出图”作为终点", level=2)
p = add_body(doc)
add_runs(p, [
    ("PISA数学素养框架把现实问题解决概括为形成（formulate）、运用（employ）、解释与评价（interpret/evaluate）三个相互往返的过程", "normal"),
    ("[9]", "sup"),
    ("。连续纹样任务中，识别重复结构并确定图形运动是“形成数学问题”，写出生成规则并交由AI执行是“运用数学”，检查接缝、对应点和形状保持情况是“解释与评价”。如果课堂止于AI输出，建模循环在最需要判断的地方中断；只有让学生依据现实要求反查模型并修正，才真正回到应用意识。", "normal")
])

add_heading(doc, "2.3　表征转换与AI认知分工", level=2)
p = add_body(doc)
add_runs(p, [
    ("杜瓦尔指出，数学理解不仅发生在同一种表征内部的处理，还依赖不同表征系统之间的转换", "normal"),
    ("[8]", "sup"),
    ("。本活动至少包含四种表征：现实纹样图像、图形运动示意、文字或符号化生成规则、AI生成图案。学生若只在“看图—出图”之间跳转，关键数学关系没有被表达；生成规则恰是连接图形运动模型与数字作品的中介。", "normal")
])
p = add_body(doc)
add_runs(p, [
    ("近期AI数学教育研究提醒，学生使用ChatGPT解决建模问题时，容易集中在前期求解而忽略结果验证与解释", "normal"),
    ("[10]", "sup"),
    ("。AI—GLM框架进一步区分AI作为生成伙伴、替代者和催化者的不同位置", "normal"),
    ("[11]", "sup"),
    ("。据此，本文确立“学生生成模型，AI执行模型；学生拥有判断权，AI输出接受检验”的认知分工。AI的价值不在替代绘制，而在快速呈现多轮结果，使规则中的含混、遗漏或矛盾变得可见。", "normal")
])

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(3)
p.paragraph_format.space_after = Pt(0)
p.add_run().add_picture(str(FIGURE), width=Cm(15.7))
add_caption(doc, "图1　生成式人工智能支持应用意识发展的教学机制")


# 3 任务与整体设计
add_heading(doc, "3　任务设计：让“生成规则”成为数学作品")
add_heading(doc, "3.1　真实任务与成果要求", level=2)
p = add_body(doc)
add_runs(p, [
    ("项目任务为：“为校园阅读空间设计一条可沿水平方向无缝延展的二方连续边饰。基本图案须含一个校园文化元素，打印单元不超过8 cm×8 cm；作品既要美观，也要让没有看过原图的人仅凭规则就能复现。”小组最终提交四项成果：初始规则卡、AI第一版输出、数学检验与修订记录、最终纹样及设计说明。成果要求把审美问题转化为可测量、可论证的数学问题。", "normal")
])
p = add_body(doc)
add_runs(p, [
    ("工具选择遵循“可检验优先”。可使用能够生成带网格图、坐标示意或可编辑矢量图的通用生成式人工智能；不要求学生编程，也不鼓励只生成写实图片。教师不上传学生姓名、肖像或未公开校园数据，并要求学生记录模型名称、输入规则和主要修订过程，落实中小学生成式人工智能使用中关于规范、安全与交叉验证的要求", "normal"),
    ("[12]", "sup"),
    ("。", "normal")
])

add_heading(doc, "3.2　“五要素生成规则”的数学结构", level=2)
p = add_body(doc)
add_runs(p, [
    ("这里的“生成规则”不是追求修辞效果的提示词，而是一份可执行的数学模型，包含五个不可缺失的要素。其最低标准是：同伴能够据此复现，AI能够据此执行，输出能够据此验收。", "normal")
])
add_caption(doc, "表2　连续纹样“五要素生成规则”")
add_table(doc,
          ["要素", "需要回答的问题", "示例表述", "常见缺陷"],
          [
              ("基本图案", "重复的最小对象是什么？边界在哪里？", "以8 cm×6 cm矩形框内的叶形线稿为一个单元", "把整条纹样误作基本图案"),
              ("运动方式", "采用平移、旋转、轴对称或其组合？", "仅作水平方向平移，保持大小和方向不变", "只写“复制”“连续排列”"),
              ("运动要素", "方向、距离、中心、角度或对称轴是什么？", "每次向右平移8 cm", "写“向右平移若干”"),
              ("重复范围", "重复多少次或覆盖怎样的区域？", "向右连续生成6个单元", "无次数、无范围"),
              ("验收条件", "怎样判定符合设计目标？", "相邻单元边界重合且无空隙、无重叠，图形不变形", "以“看起来连续”为标准"),
          ],
          [2.0, 4.0, 5.3, 4.4], font_size=8.5)

p = add_body(doc)
add_runs(p, [
    ("以最简单的平移纹样为例，若单元宽度为d，同一对应点在第k个单元中的位置可表示为“横向位置增加kd，纵向位置不变”。对七年级学生不必引入向量符号，但应要求他们借助网格或刻度说明：所有对应点移动方向一致、距离相等，图形的形状和大小保持不变。由此，规则的正确性由可测的对应关系保证，而不是由视觉相似保证。", "normal")
])

add_heading(doc, "3.3　六课时活动安排", level=2)
add_caption(doc, "表3　“连续纹样”综合与实践的总体安排")
add_table(doc,
          ["课时", "核心问题", "主要活动", "AI位置", "阶段成果"],
          [
              ("课外1", "生活中的连续纹样怎样“连续”？", "采集、裁切、分类纹样，尝试圈出重复单元", "不使用AI，保留原始观察", "纹样素材卡"),
              ("课内1", "一个复杂纹样可由什么运动形成？", "比较实例，识别基本图案及平移、旋转、轴对称", "用于展示不同解释，不直接判定", "运动分析图"),
              ("课内2", "怎样让别人只凭规则复现？", "把自然语言改写为五要素规则，同伴盲绘检验", "作为第二执行者", "规则卡1.0"),
              ("课内3", "AI画出来就说明规则正确吗？", "生成带网格图，测对应点、查边界、记录偏差", "执行并暴露歧义", "检验单与规则卡2.0"),
              ("课内4", "规则怎样适用于新图案？", "比较多种规则，完成组合变换与迁移挑战", "批量呈现方案供比较", "最终作品与说明"),
              ("课外2", "作品怎样服务校园空间？", "按尺寸排版，接受同伴验收，反思人机分工", "辅助格式整理", "项目档案袋"),
          ],
          [1.5, 4.2, 5.1, 3.1, 2.7], font_size=8.2)


# 4 核心课时
add_heading(doc, "4　核心课时实施：从含混指令到可检验规则")
add_heading(doc, "4.1　制造认知冲突：一幅“像连续纹样”的错误作品", level=2)
p = add_body(doc)
add_runs(p, [
    ("教师先将各组未经修改的自然语言交给AI。某组输入：“把这个叶形图案向右连续复制，做成漂亮的校园边饰。”AI输出了六个叶形，但相邻间距不一，其中一个叶形方向改变。教师不直接评价，只追问：“它完成任务了吗？你依据什么判断？”学生容易出现两种回答：“整体看起来是连续的”和“第三个图形翻转了，所以不对”。教师据此把审美判断推进为数学判断：平移前后的对应点发生了怎样的变化？形状和方向应不应该改变？每次移动的距离是否相同？", "normal")
])
p = add_body(doc)
add_runs(p, [
    ("预设课堂片段如下。", "normal")
])
dialogues = [
    ("师：", "如果只告诉AI“向右连续复制”，还缺了什么？"),
    ("生1：", "缺平移多少。它可以每次移动不同的距离。"),
    ("师：", "“移动8厘米”够了吗？"),
    ("生2：", "还要说谁移动、向哪个方向，以及移动后形状和方向保持不变。"),
    ("师：", "怎样让“无缝”也能检查？"),
    ("生3：", "加网格或标出单元边界，检查相邻边界是否接上，有没有空隙和重叠。")
]
for who, words in dialogues:
    p = add_body(doc, indent=False, line=1.25, after=0)
    p.paragraph_format.left_indent = Cm(0.8)
    r = p.add_run(who)
    set_run_font(r, cn=HEAD_CN, size=9.8, bold=True, color=BLUE)
    r = p.add_run(words)
    set_run_font(r, size=9.8)

add_heading(doc, "4.2　同伴“盲执行”：先检验数学表达，再交给AI", level=2)
p = add_body(doc)
add_runs(p, [
    ("直接与AI对话容易使学生把输出偏差归因于“AI不听话”。为突出数学表达，教师先安排同伴盲执行：A组只交规则，不给原图；B组按规则在方格纸上画出前三个单元，并标注无法确定之处。例如，“将基本图案关于竖直直线对称后向右移动”至少存在对称轴位置、先后顺序和平移距离三处不确定。同伴的追问成为规则缺口的第一层证据。学生据此补全五要素，再由AI作为第二个执行者。人与AI对同一规则产生不同结果时，讨论焦点不是“谁更聪明”，而是哪一条规则仍不能排除多种解释。", "normal")
])

add_heading(doc, "4.3　建立检验链：局部运动正确不等于整体能够延展", level=2)
p = add_body(doc)
add_runs(p, [
    ("AI生成后，各组不能以“满意/不满意”作结，而须完成三步检验。第一，检验运动：任选不少于两个对应点，比较其方向和距离；若使用旋转或轴对称，还要核对中心、角度或对称轴。第二，检验不变量：平移和旋转前后图形全等，轴对称前后形状、大小不变，不能接受AI擅自增删花瓣、改变比例或润饰轮廓。第三，检验连续性：将一个重复单元平移到相邻位置，观察边界是否重合，继续延展后是否产生空隙、重叠或节奏突变。", "normal")
])
p = add_body(doc)
add_runs(p, [
    ("这一环节刻意设置“局部成立但整体失败”的情形。例如，每个叶形确实都由前一个叶形向右平移8 cm得到，但若基本单元宽度为7.5 cm，拼接后会出现0.5 cm空隙。学生由此认识：某一次图形运动正确，只能说明局部对应关系成立；要形成连续纹样，还须使运动距离与重复单元边界相匹配。该辨析把简单的操作活动提升为“条件是否充分”的数学讨论。", "normal")
])

add_caption(doc, "表4　AI输出的数学检验记录单（节选）")
add_table(doc,
          ["检验维度", "检验问题", "可接受的证据", "可能的修正规则"],
          [
              ("对应点", "各单元对应点的方向、距离是否一致？", "网格计数、刻度测量、重叠比较", "补充方向和精确距离"),
              ("图形不变量", "形状、大小、方向是否按运动方式保持？", "描图纸覆盖或关键长度、角度比较", "增加“不得变形、增删或改向”"),
              ("单元边界", "相邻边界能否吻合？", "标出矩形单元框，比较左右边界", "令平移距离等于单元宽度"),
              ("整体延展", "继续重复后是否无缝、无重叠、无突变？", "至少检查三个连续单元", "明确重复次数和验收条件"),
              ("规则一致", "输出能否逐条对应五要素？", "在图上标注每条规则的落点", "删除装饰性含混词，补足数学约束"),
          ],
          [2.2, 4.4, 4.8, 4.7], font_size=8.4)

add_heading(doc, "4.4　修正规则与迁移：从“这一幅对了”到“这一类可生成”", level=2)
p = add_body(doc)
add_runs(p, [
    ("修订时，学生须采用“偏差—原因—修改—再检验”的句式记录。例如：“第二个单元方向改变；原规则未说明只作平移且方向保持；补充‘基本图案不作旋转或镜像’；重新生成后用三组对应点验证。”这样，AI的错误不再只是需要重试的技术故障，而成为反查模型与语言精度的学习资源。", "normal")
])
p = add_body(doc)
add_runs(p, [
    ("迁移任务不再要求重复计算，而是改变条件：“现将相邻单元改为一正一反，并仍要水平无缝延展，原有规则哪些保留、哪些必须改变？”学生需要在“先轴对称再平移”与“先平移再轴对称”之间辨析操作顺序，并说明对称轴位置与平移距离。最终交流不评选“最漂亮的图”，而比较不同小组规则的简洁性、唯一性和可推广性。", "normal")
])


# 5 评价
add_heading(doc, "5　评价设计：捕捉应用意识发展的过程证据")
add_heading(doc, "5.1　证据链与评价量规", level=2)
p = add_body(doc)
add_runs(p, [
    ("应用意识不能由一幅最终作品直接推断。本文把评价证据安排在学习链的关键节点：最初的纹样圈画反映结构提取；规则卡1.0与2.0的差异反映模型表达的精确化；AI第一版与修订版的标注反映检验意识；迁移任务反映学生能否把方法用于新情境。教师同时收集小组对话、同伴盲执行中的疑问和个人反思，避免只以小组作品代表每名学生。", "normal")
])
add_caption(doc, "表5　项目学习评价量规")
add_table(doc,
          ["评价维度及权重", "达成", "进阶", "优秀"],
          [
              ("结构提取（20%）", "能圈出一个可重复单元", "能说明基本图案与整体纹样的关系", "能比较不同单元选取并说明优劣"),
              ("规则表达（30%）", "包含运动方式和大致方向", "五要素基本完整，同伴可以复现", "规则简洁、唯一，能用图示或符号互证"),
              ("数学检验（30%）", "能指出明显空隙或变形", "能用对应点、距离或边界给出证据", "能区分局部与整体，并据证据追溯规则缺口"),
              ("修正迁移（20%）", "能按提示修改规则", "能独立完成“偏差—原因—修改—复检”", "能迁移到组合运动并解释适用条件"),
          ],
          [3.1, 4.3, 4.3, 4.3], font_size=8.5)

add_heading(doc, "5.2　拟采用的分析指标", level=2)
p = add_body(doc)
add_runs(p, [
    ("正式实施时，不宜只报告“学生喜欢AI”或作品数量，而应分析规则和证据的质量。可从四方面编码：一是规则完整度，统计五要素的出现与准确情况；二是表征一致性，比较学生的图形标注、语言规则与AI产物能否相互对应；三是证据层级，区分“凭视觉判断”“指出数学关系”“用测量或重叠验证”“依据验证结果反改模型”；四是迁移水平，判断学生能否在基本图案或运动方式变化后重新建模。若条件允许，可采用前后同构任务比较同一学生规则卡的变化，并以代表性学生作品作质性分析。", "normal")
])


# 6 讨论
add_heading(doc, "6　讨论：AI何以支持、又在何处退出")
add_heading(doc, "6.1　AI的独特价值在于“外显规则后果”", level=2)
p = add_body(doc)
add_runs(p, [
    ("传统剪拼、描图和电子白板同样能够呈现图形运动，AI并非不可替代。它在本设计中的相对优势是：对同一规则迅速生成多轮结果，把自然语言中的遗漏和多义性转化为可比较的视觉后果；同时可以按学生提出的新约束即时重画，使“提出模型—检验—修正”在一节课内多次发生。也正因如此，论文的贡献不是证明AI比传统工具更好，而是说明在什么认知分工下，AI的生成能力能够服务于数学建模而不吞噬建模。", "normal")
])

add_heading(doc, "6.2　与同题研究的差异：从技术操作走向规则可检验", level=2)
p = add_body(doc)
add_runs(p, [
    ("与既有平面镶嵌项目相比，本文保留真实设计、跨学科审美与作品评价，但把研究对象缩小为“规则如何形成和修正”；与一般GAI教学设计相比，本文不以资源生成或个性化答疑为主，而把学生书写的数学规则作为中介产物。因而，可检验的创新点不是纹样题材，也不是使用某个模型，而是三项机制：同伴与AI的双重执行、基于对应点与边界的证据链、由输出偏差反查规则缺口的迭代记录。", "normal")
])

add_heading(doc, "6.3　实施边界与风险", level=2)
p = add_body(doc)
add_runs(p, [
    ("第一，AI图像生成具有随机性，教师应保留规则和版本记录，并准备方格纸或动态几何软件作为替代执行工具；若模型不能稳定输出网格或矢量图，就不适合承担精确验证。第二，提示词不能变成新的“技巧教学”，评价应指向数学要素而非词藻长度。第三，小组共用设备可能遮蔽个体差异，核心判断与迁移任务应个人完成。第四，AI给出的图案不构成数学证明，任何“无缝”“全等”或“保持不变”的结论都必须回到图形关系和测量证据。第五，学生作品涉及著作权和数据安全时，应使用自绘基本图案，避免上传人物肖像和受保护素材。", "normal")
])


# 7 结语
add_heading(doc, "7　结语")
p = add_body(doc)
add_runs(p, [
    ("生成式人工智能进入初中数学课堂，真正值得研究的不是它能生成多少资源，而是它改变了哪些数学活动。在连续纹样项目中，让学生把图形运动写成五要素生成规则，再让AI执行、暴露歧义并接受数学检验，可以使“用数学解释和设计现实事物”从口号落到具体行为。由“图形运动”到“生成规则”的转换，也使学生留下可比较的模型、证据和修订痕迹。下一步需在真实课堂中依据规则卡、检验单、个人迁移任务及访谈材料，考察这一机制对不同基础学生的适切性，并据实报告其效果与限制。", "normal")
])


# 参考文献
add_heading(doc, "参考文献")
refs = [
    "[1] 中华人民共和国教育部. 义务教育数学课程标准（2022年版）[S]. 北京：北京师范大学出版社，2022.",
    "[2] 李大潜. 义务教育教科书（五·四学制）数学教学参考资料：七年级上册[M]. 上海：上海教育出版社，2024.",
    "[3] 郭衍，曹一鸣. 综合与实践：从主题活动到项目学习[J]. 数学教育学报，2022，31（5）：9-13.",
    "[4] 李佩乔，杨雪花，张海湘. 指向应用意识的初中数学项目式学习的设计路径研究——以“人教版初中数学八年级平面镶嵌问题”为例[J]. 理科考试研究，2025，32（2）：2-5.",
    "[5] 苏国东. 智能技术支持下的初中数学深度学习实践研究——以“平面图形的镶嵌”为例[J]. 数学之友，2025（17）：93-96.",
    "[6] 赵文君，李星，唐海军. 生成式人工智能赋能数学概念教学的案例研究——以“相似多边形”为例[J]. 福建教育，2026（4）：119-124.",
    "[7] 沈爱桐，杨亚平. 人工智能驱动初中数学“综合与实践”活动的实践路径[J]. 教学与管理，2026（3）：65-70.",
    "[8] DUVAL R. A cognitive analysis of problems of comprehension in a learning of mathematics[J]. Educational Studies in Mathematics, 2006, 61: 103-131.",
    "[9] OECD. PISA 2022 Assessment and Analytical Framework[M/OL]. Paris: OECD Publishing, 2023.",
    "[10] BULUT N, BORROMEO FERRI R. AI-powered learning: exploring secondary school students’ use of ChatGPT for solving mathematical modelling problems[J]. ZDM-Mathematics Education, 2026, 58: 993-1009.",
    "[11] CHO H. Relocating the locus of generation: the AI-GLM framework for generative learning in AI-mediated mathematics education[J/OL]. Educational Psychology Review, 2026. DOI:10.1007/s10648-026-10210-y.",
    "[12] 教育部基础教育教学指导委员会. 中小学生成式人工智能使用指南（2025年版）[S]. 2025.",
]
for ref in refs:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(0.74)
    p.paragraph_format.first_line_indent = Cm(-0.74)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(ref)
    set_run_font(r, size=8.8)
    set_widow_control(p)

doc.save(OUT)
print(OUT)
